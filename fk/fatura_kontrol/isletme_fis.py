"""İşletme / serbest meslek defteri: faturalardan Luca "Excel'den fiş aktarma" dosyası.

Kullanıcının gönderdiği örnek Excel (ÖRNEK FİŞ EXCEL.xlsx) düzeninde, fatura başına (KDV oranı başına) bir satır:
İŞLEM · KATEGORİ · BELGE TÜRÜ · EVRAK TARİHİ · KAYIT TARİHİ · SERİ NO · EVRAK NO · TCKN/VKN · VERGİ DAİRESİ ·
SOYADI ÜNVAN · ADI DEVAMI · ADRES · CARİ HESAP · KDV İSTİSNASI · KOD · BELGE TÜRÜ(DB) · ALIŞ/SATIŞ TÜRÜ ·
KAYIT ALT TÜRÜ · MAL VE HİZMET KODU · AÇIKLAMA · MİKTAR · B.FİYAT · TUTAR · TEVKİFAT · KDV ORANI · KDV TUTARI ·
TOPLAM TUTAR · KREDİLİ TUTAR · STOPAJ KODU · STOPAJ TUTARI · DÖNEMSELLİK İLKESİ · FAALIYET KODU · ÖDEME TÜRÜ

Kurallar (kullanıcı kararları, 22.10.2026):
  * alış → İşlem "Gider", belge türü "Diğer Alışlar"; satış → "Gelir", "Diğer Satışlar" (satırda değiştirilebilir)
  * fatura başına tek satır; aynı KDV oranındaki kalemler birleşir, farklı oranlar ayrı satır
  * ünvan faturadan, kısaltmalar açılarak (A.Ş. → ANONİM ŞİRKETİ …); ilk 20 karakter SOYADI ÜNVAN, kalanı ADI DEVAMI;
    gerçek kişide soyadı / adı
  * kayıt alt türü, KDV istisnası, stopaj kodu vb. Luca'daki listelerden; kullanıcının seçimleri karşı tarafa göre öğrenilir
"""
from __future__ import annotations

import json
import re
import xml.etree.ElementTree as ET
from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from . import depo
from .fatura_okuyucu import NS, _hamlar, xml_oku

BASLIKLAR = ["İŞLEM", "KATEGORİ", "BELGE TÜRÜ", "EVRAK TARİHİ", "KAYIT TARİHİ", "SERİ NO", "EVRAK NO", "TCKN/VKN",
             "VERGİ DAİRESİ", "SOYADI ÜNVAN", "ADI DEVAMI", "ADRES", "CARİ HESAP", "KDV İSTİSNASI", "KOD", "BELGE TÜRÜ(DB)",
             "ALIŞ/SATIŞ TÜRÜ", "KAYIT ALT TÜRÜ", "MAL VE HİZMET KODU", "AÇIKLAMA", "MİKTAR", "B.FİYAT", "TUTAR", "TEVKİFAT",
             "KDV ORANI", "KDV TUTARI", "TOPLAM TUTAR", "KREDİLİ TUTAR", "STOPAJ KODU", "STOPAJ TUTARI", "DÖNEMSELLİK İLKESİ",
             "FAALIYET KODU", "ÖDEME TÜRÜ"]
ANAHTAR = ["islem", "kategori", "belge_turu", "evrak_tarihi", "kayit_tarihi", "seri_no", "evrak_no", "vkn", "vergi_dairesi",
           "unvan1", "unvan2", "adres", "cari_hesap", "kdv_istisnasi", "kod", "belge_turu_db", "alis_satis_turu", "kayit_alt_turu",
           "mal_hizmet_kodu", "aciklama", "miktar", "birim_fiyat", "tutar", "tevkifat", "kdv_orani", "kdv_tutari", "toplam",
           "kredili", "stopaj_kodu", "stopaj_tutari", "donemsellik", "faaliyet_kodu", "odeme_turu"]
SAYISAL = {"miktar", "birim_fiyat", "tutar", "kdv_orani", "kdv_tutari", "toplam", "kredili", "stopaj_tutari"}
# kullanıcının satırda seçtiği ve karşı tarafa göre hatırlanan alanlar
OGRENILEN = ["belge_turu", "alis_satis_turu", "kayit_alt_turu", "kdv_istisnasi", "stopaj_kodu", "cari_hesap", "odeme_turu"]
KATEGORILER = ["Defter Fişleri", "Serbest Meslek Fişleri"]
# Luca'dan listeler alınana kadar önerilen değerler
VARSAYILAN_LISTELER = {
    "islem": ["Gelir", "Gider"],
    "kategori": KATEGORILER,
    "belge_turu": ["Diğer Alışlar", "Diğer Satışlar", "Satış", "Alıştan İade", "Satıştan İade"],
    "belge_turu_db": ["e-Fatura", "e-Arşiv Fatura", "e-Serbest Meslek Makbuzu", "e-Bilet", "Kağıt Fatura"],
    "alis_satis_turu": ["Normal Alım", "Normal Satış"],
    "tevkifat": ["Yok"],
}
AYAR_JSON = "isletme_fis_ayar.json"
SECIM_JSON = "isletme_fis_secimleri.json"
LISTE_JSON = "luca_isletme_listeleri.json"

KISALTMALAR = [   # (desen, açılımı) — büyük harfli metinde, sözcük sınırında
    (r"A\.?\s?Ş\.?", "ANONİM ŞİRKETİ"),
    (r"LTD\.?\s?ŞTİ\.?", "LİMİTED ŞİRKETİ"),
    (r"LİMİTED\s+ŞTİ\.?", "LİMİTED ŞİRKETİ"),
    (r"ŞTİ\.?", "ŞİRKETİ"), (r"LTD\.?", "LİMİTED"),
    (r"TİC\.?", "TİCARET"), (r"SAN\.?", "SANAYİ"), (r"İNŞ\.?", "İNŞAAT"), (r"PAZ\.?", "PAZARLAMA"),
    (r"HİZM?\.?", "HİZMETLERİ"), (r"MÜH\.?", "MÜHENDİSLİK"), (r"TUR\.?", "TURİZM"), (r"NAK\.?", "NAKLİYAT"),
    (r"GIDA\.?", "GIDA"), (r"OTOM\.?", "OTOMOTİV"), (r"ELEKT?\.?", "ELEKTRİK"), (r"DIŞ\.?", "DIŞ"),
    (r"İTH\.?", "İTHALAT"), (r"İHR\.?", "İHRACAT"), (r"KOLL\.?", "KOLLEKTİF"), (r"KOM\.?", "KOMANDİT"),
]
IL_PLAKA = {"ADANA": 1, "ADIYAMAN": 2, "AFYONKARAHİSAR": 3, "AFYON": 3, "AĞRI": 4, "AMASYA": 5, "ANKARA": 6, "ANTALYA": 7,
            "ARTVİN": 8, "AYDIN": 9, "BALIKESİR": 10, "BİLECİK": 11, "BİNGÖL": 12, "BİTLİS": 13, "BOLU": 14, "BURDUR": 15,
            "BURSA": 16, "ÇANAKKALE": 17, "ÇANKIRI": 18, "ÇORUM": 19, "DENİZLİ": 20, "DİYARBAKIR": 21, "EDİRNE": 22,
            "ELAZIĞ": 23, "ERZİNCAN": 24, "ERZURUM": 25, "ESKİŞEHİR": 26, "GAZİANTEP": 27, "GİRESUN": 28, "GÜMÜŞHANE": 29,
            "HAKKARİ": 30, "HATAY": 31, "ISPARTA": 32, "MERSİN": 33, "İÇEL": 33, "İSTANBUL": 34, "İZMİR": 35, "KARS": 36,
            "KASTAMONU": 37, "KAYSERİ": 38, "KIRKLARELİ": 39, "KIRŞEHİR": 40, "KOCAELİ": 41, "KONYA": 42, "KÜTAHYA": 43,
            "MALATYA": 44, "MANİSA": 45, "KAHRAMANMARAŞ": 46, "MARDİN": 47, "MUĞLA": 48, "MUŞ": 49, "NEVŞEHİR": 50,
            "NİĞDE": 51, "ORDU": 52, "RİZE": 53, "SAKARYA": 54, "SAMSUN": 55, "SİİRT": 56, "SİNOP": 57, "SİVAS": 58,
            "TEKİRDAĞ": 59, "TOKAT": 60, "TRABZON": 61, "TUNCELİ": 62, "ŞANLIURFA": 63, "UŞAK": 64, "VAN": 65, "YOZGAT": 66,
            "ZONGULDAK": 67, "AKSARAY": 68, "BAYBURT": 69, "KARAMAN": 70, "KIRIKKALE": 71, "BATMAN": 72, "ŞIRNAK": 73,
            "BARTIN": 74, "ARDAHAN": 75, "IĞDIR": 76, "YALOVA": 77, "KARABÜK": 78, "KİLİS": 79, "OSMANİYE": 80, "DÜZCE": 81}
# Luca'daki vergi dairesi kodları (il plakası + daire no). Luca'dan liste alınınca oradakiler önce gelir.
VD_KODLARI = {"KADIKÖY|34": "034272", "SARAYLAR|20": "020201"}


def buyuk(s: str) -> str:
    return (s or "").replace("i", "İ").replace("ı", "I").upper()


def sade(s: str) -> str:
    return re.sub(r"\s+", " ", buyuk(s)).strip()


def kisaltmalari_ac(unvan: str) -> str:
    s = sade(unvan)
    for desen, acilim in KISALTMALAR:
        s = re.sub(rf"(?<![A-ZÇĞİÖŞÜ]){desen}(?![A-ZÇĞİÖŞÜ])", acilim, s)
    return re.sub(r"\s+", " ", s).strip()


def unvan_bol(unvan: str, sinir: int = 20) -> tuple[str, str]:
    """Örnek Excel'deki gibi sözcük gözetmeden: ilk 20 karakter SOYADI ÜNVAN, kalanı ADI DEVAMI."""
    return unvan[:sinir], unvan[sinir:]


def _t(el, yol):
    if el is None:
        return ""
    x = el.find(yol, NS)
    return (x.text or "").strip() if x is not None and x.text else ""


def taraf_ayrinti(root, yol: str) -> dict:
    """Faturadaki taraf: kimlik, ünvan / ad-soyad, vergi dairesi, adres, il."""
    p = root.find(yol + "/cac:Party", NS)
    if p is None:
        return {}
    vkn = ""
    for i in p.findall("cac:PartyIdentification/cbc:ID", NS):
        if i.get("schemeID") in ("VKN", "TCKN") and i.text:
            vkn = i.text.strip()
            break
    a = p.find("cac:PostalAddress", NS)
    sokak = " ".join(x for x in [_t(a, "cbc:StreetName"), _t(a, "cbc:BuildingName"),
                                  ("No:" + _t(a, "cbc:BuildingNumber")) if _t(a, "cbc:BuildingNumber") else "",
                                  _t(a, "cbc:Room")] if x)
    ilce, il = _t(a, "cbc:CitySubdivisionName"), _t(a, "cbc:CityName")
    adres = (sokak + (" " if sokak and (ilce or il) else "") + "/".join(x for x in [ilce, il] if x)).strip()
    return {"vkn": vkn, "unvan": _t(p, "cac:PartyName/cbc:Name"), "ad": _t(p, "cac:Person/cbc:FirstName"),
            "soyad": _t(p, "cac:Person/cbc:FamilyName"), "vd": _t(p, "cac:PartyTaxScheme/cac:TaxScheme/cbc:Name"),
            "adres": adres, "il": il}


def kalem_miktarlari(root) -> dict[int, Decimal]:
    """KDV oranı → kalemlerin miktar toplamı."""
    out: dict[int, Decimal] = {}
    for ln in root.findall("cac:InvoiceLine", NS):
        try:
            m = Decimal(_t(ln, "cbc:InvoicedQuantity") or "0")
        except Exception:
            m = Decimal("0")
        o = _t(ln, "cac:TaxTotal/cac:TaxSubtotal/cbc:Percent") or _t(ln, "cac:TaxTotal/cac:TaxSubtotal/cac:TaxCategory/cbc:Percent")
        try:
            oran = int(Decimal(o or "0"))
        except Exception:
            oran = 0
        out[oran] = out.get(oran, Decimal("0")) + m
    return out


def vergi_dairesi(ad: str, il: str, listeler: dict | None = None) -> str:
    """'034272  KADIKÖY VERGİ DAİRESİ' biçimi (Luca). Kod bilinmiyorsa yalnız ad."""
    a = sade(ad)
    a = re.sub(r"\s*(VERGİ DAİRESİ( MÜDÜRLÜĞÜ)?|V\.?\s?D\.?)$", "", a).strip()
    if not a:
        return ""
    plaka = IL_PLAKA.get(sade(il).split("/")[-1].strip(), 0)
    tam = f"{a} VERGİ DAİRESİ"
    for s in (listeler or {}).get("vergi_dairesi", []):        # Luca'dan alınan liste: "034272  KADIKÖY VERGİ DAİRESİ"
        m = re.match(r"^(\d{6})\s+(.+)$", str(s).strip())
        if m and sade(m.group(2)) in (tam, a) and (not plaka or int(m.group(1)[:3]) == plaka):
            return str(s).strip()
    kod = VD_KODLARI.get(f"{a}|{plaka}")
    return f"{kod}  {tam}" if kod else tam


def _tarih(d: date) -> str:
    return d.strftime("%d/%m/%Y")


def _r(x) -> float:
    return float(Decimal(x).quantize(Decimal("0.01"), ROUND_HALF_UP))


def _json_oku(y: Path, bos):
    try:
        return json.loads(y.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return bos


def ayar_oku(firma_ad: str) -> dict:
    v = _json_oku(depo.firma_klasoru(firma_ad) / AYAR_JSON, {})
    return {"kategori": v.get("kategori") or KATEGORILER[0], "faaliyet_kodu": v.get("faaliyet_kodu") or "",
            "kimlikler": v.get("kimlikler") or [], "kayit_alt_turu_gider": v.get("kayit_alt_turu_gider") or "",
            "kayit_alt_turu_gelir": v.get("kayit_alt_turu_gelir") or ""}


def ayar_yaz(firma_ad: str, ayar: dict) -> dict:
    k = depo.firma_klasoru(firma_ad); k.mkdir(parents=True, exist_ok=True)
    eski = ayar_oku(firma_ad)
    eski.update({x: ayar[x] for x in ("kategori", "faaliyet_kodu", "kimlikler", "kayit_alt_turu_gider", "kayit_alt_turu_gelir") if x in ayar})
    eski["kimlikler"] = sorted({re.sub(r"\D", "", str(x)) for x in eski["kimlikler"] if re.sub(r"\D", "", str(x))})
    (k / AYAR_JSON).write_text(json.dumps(eski, ensure_ascii=False, indent=1), encoding="utf-8")
    return eski


def secimler_oku(firma_ad: str) -> dict:
    return _json_oku(depo.firma_klasoru(firma_ad) / SECIM_JSON, {})


def secimleri_ogren(firma_ad: str, satirlar: list[dict]):
    """Excel'e aktarılan satırlardaki seçimler karşı tarafa (VKN + işlem) göre hatırlanır."""
    k = depo.firma_klasoru(firma_ad); k.mkdir(parents=True, exist_ok=True)
    s = secimler_oku(firma_ad)
    for r in satirlar:
        if not r.get("vkn"):
            continue
        anahtar = f"{r['vkn']}|{r.get('islem', '')}"
        s[anahtar] = {x: r.get(x, "") for x in OGRENILEN}
    (k / SECIM_JSON).write_text(json.dumps(s, ensure_ascii=False, indent=1), encoding="utf-8")


def listeler_oku() -> dict:
    """Luca'dan alınmış açılır listeler (MuhaWin › Luca penceresi) + önerilen değerler."""
    v = _json_oku(depo.veri() / LISTE_JSON, {})
    out = {k: list(x) for k, x in VARSAYILAN_LISTELER.items()}
    for k, x in (v.get("alanlar") or {}).items():
        out[k] = [str(y) for y in x if str(y).strip()] or out.get(k, [])
    out["_kaynak"] = {"luca": bool(v.get("alanlar")), "zaman": v.get("zaman", ""), "ham": len(v.get("ham") or [])}
    return out


# Luca ekranındaki <select> → alan eşlemesi (ad / id / etiket metnine göre)
ALAN_DESENLERI = [
    ("kayit_alt_turu", r"alt.?t[üu]r|altTur|ALT_TUR"),
    ("alis_satis_turu", r"al[ıi][şs].?sat[ıi][şs].?t[üu]r|ALIS_SATIS|alisSatis"),
    ("belge_turu_db", r"belge.?t[üu]r.*\(?db|e-?belge|ebelge|BELGE_TURU_DB"),
    ("belge_turu", r"belge.?t[üu]r|BELGE_TUR"),
    ("kdv_istisnasi", r"istisna"),
    ("stopaj_kodu", r"stopaj"),
    ("tevkifat", r"tevkif"),
    ("vergi_dairesi", r"vergi.?daire|VERGI_DAIRE|vd"),
    ("faaliyet_kodu", r"faaliyet|nace"),
    ("odeme_turu", r"[öo]deme"),
    ("kategori", r"kategori"),
    ("donemsellik", r"d[öo]nemsel"),
]


def listeleri_kaydet(ham: list[dict], zaman: str = "") -> dict:
    """ham: [{ad, id, etiket, secenekler: [metin…]}] — MuhaWin'in Luca penceresinden topladığı açılır listeler."""
    alanlar: dict[str, list[str]] = {}
    for s in ham:
        metin = " ".join(str(s.get(x) or "") for x in ("etiket", "ad", "id"))
        for alan, desen in ALAN_DESENLERI:
            if re.search(desen, metin, re.I):
                L = [str(o).strip() for o in s.get("secenekler") or [] if str(o).strip() and not re.match(r"^(seçiniz|seciniz|-+)$", str(o).strip(), re.I)]
                if L and len(L) > len(alanlar.get(alan, [])):
                    alanlar[alan] = L
                break
    eski = _json_oku(depo.veri() / LISTE_JSON, {})
    yeni_alanlar = {**(eski.get("alanlar") or {}), **alanlar}
    v = {"zaman": zaman, "alanlar": yeni_alanlar, "ham": ham}
    (depo.veri() / LISTE_JSON).parent.mkdir(parents=True, exist_ok=True)
    (depo.veri() / LISTE_JSON).write_text(json.dumps(v, ensure_ascii=False, indent=1), encoding="utf-8")
    return {"bulunan": sorted(alanlar), "liste_sayisi": len(ham)}


def satirlar_hazirla(zipler: list[Path], firma_ad: str, kimlikler: set[str], bas: date | None = None,
                     bit: date | None = None) -> dict:
    """Faturalardan fiş satırları. kimlikler: firmanın VKN / TCKN'leri (şahıs firmasında ikisi de olabilir)."""
    ayar = ayar_oku(firma_ad)
    kimlikler = {k for k in kimlikler | set(ayar["kimlikler"]) if k}
    sec = secimler_oku(firma_ad)
    listeler = listeler_oku()
    satirlar, disarida, gorulen = [], [], set()
    for veri, kaynak in _hamlar(zipler):
        f = xml_oku(veri, kaynak)
        if f is None:
            continue
        anahtar = f.ettn or (f.satici_vkn, f.belge_no)
        if anahtar in gorulen:
            continue
        gorulen.add(anahtar)
        if bas and bit and not (bas <= f.tarih <= bit):
            continue
        if f.alici_vkn in kimlikler:
            yon = "ALIS"
        elif f.satici_vkn in kimlikler:
            yon = "SATIS"
        else:
            disarida.append({"belge_no": f.belge_no, "tarih": f.tarih.isoformat(), "satici": f.satici_unvan, "satici_vkn": f.satici_vkn,
                             "alici": f.alici_unvan, "alici_vkn": f.alici_vkn})
            continue
        try:
            root = ET.fromstring(veri)
        except ET.ParseError:
            root = None
        karsi = (taraf_ayrinti(root, "cac:AccountingSupplierParty" if yon == "ALIS" else "cac:AccountingCustomerParty")
                 if root is not None and root.tag.endswith("Invoice") else {})
        karsi_vkn = f.satici_vkn if yon == "ALIS" else f.alici_vkn
        karsi_ad = f.satici_unvan if yon == "ALIS" else f.alici_unvan
        iade = (f.tip or "").upper() == "IADE"
        # alış yönünde iade faturası: müşterimiz bize iade etti → bizim satıştan iademiz (gider); satış yönünde iade: alıştan iade (gelir)
        if yon == "ALIS" and iade:
            islem, belge = "Gider", "Satıştan İade"
        elif yon == "SATIS" and iade:
            islem, belge = "Gelir", "Alıştan İade"
        else:
            islem = "Gider" if yon == "ALIS" else "Gelir"
            belge = "Diğer Alışlar" if yon == "ALIS" else "Diğer Satışlar"
        if len(karsi_vkn) == 11 and (karsi.get("soyad") or karsi.get("ad")):
            u1, u2 = sade(karsi.get("soyad", "")), sade(karsi.get("ad", ""))
        else:
            u1, u2 = unvan_bol(kisaltmalari_ac(karsi.get("unvan") or karsi_ad))
        profil = (f.profil or "").upper()
        db = "e-Serbest Meslek Makbuzu" if "SERBEST" in (f.satici_unvan or "").upper() and not profil else \
            "e-Arşiv Fatura" if "EARSIV" in profil else "e-Fatura"
        miktarlar = kalem_miktarlari(root) if root is not None and root.tag.endswith("Invoice") else {}
        onceki = sec.get(f"{karsi_vkn}|{islem}") or {}
        oranlar = f.matrah_oranlari or {Decimal("0"): f.matrah}
        tevk = {d["oran"] for d in f.tevkifat_detay if d.get("oran")}
        for oran, matrah in sorted(oranlar.items(), key=lambda x: -x[0]):
            o = int(oran)
            kdv = f.kdv_oranlari.get(oran, Decimal("0"))
            mik = miktarlar.get(o) or Decimal("1")
            r = {k: "" for k in ANAHTAR}
            r.update({
                "islem": islem, "kategori": ayar["kategori"], "belge_turu": belge,
                "evrak_tarihi": _tarih(f.tarih), "kayit_tarihi": _tarih(f.tarih), "seri_no": "", "evrak_no": f.belge_no,
                "vkn": karsi_vkn, "vergi_dairesi": vergi_dairesi(karsi.get("vd", ""), karsi.get("il", ""), listeler),
                "unvan1": u1, "unvan2": u2, "adres": karsi.get("adres", ""), "belge_turu_db": db,
                "alis_satis_turu": "Normal Alım" if islem == "Gider" else "Normal Satış",
                "kayit_alt_turu": ayar["kayit_alt_turu_gider" if islem == "Gider" else "kayit_alt_turu_gelir"],
                "mal_hizmet_kodu": " ", "aciklama": karsi_ad, "miktar": float(mik),
                "birim_fiyat": _r(matrah / mik) if mik else _r(matrah), "tutar": _r(matrah),
                "tevkifat": ", ".join(sorted(tevk)) if tevk else "Yok", "kdv_orani": o, "kdv_tutari": _r(kdv),
                "toplam": _r(matrah + kdv), "kredili": 0, "stopaj_tutari": 0, "faaliyet_kodu": ayar["faaliyet_kodu"],
                "kdv_istisnasi": (f.istisna or "") if not kdv else "",
            })
            for k in OGRENILEN:
                if onceki.get(k):
                    r[k] = onceki[k]
            r["_belge"] = f.belge_no
            r["_yon"] = yon
            r["_ogrenildi"] = bool(onceki)
            satirlar.append(r)
    satirlar.sort(key=lambda r: (r["islem"], r["evrak_tarihi"][6:] + r["evrak_tarihi"][3:5] + r["evrak_tarihi"][:2], r["evrak_no"]))
    return {"satirlar": satirlar, "disarida": disarida, "ayar": ayar, "kimlikler": sorted(kimlikler)}


def excel_yaz(satirlar: list[dict], yol: Path) -> Path:
    """Örnek Excel ile aynı düzen: tek sayfa, başlık satırı, tarihler gg/aa/yyyy metin, tutarlar sayı (#,##0.00)."""
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "-"
    ws.append(BASLIKLAR)
    for r in satirlar:
        ws.append([(r.get(k) if k in SAYISAL else ("" if r.get(k) is None else str(r.get(k)))) for k in ANAHTAR])
    for satir in ws.iter_rows(min_row=2):
        for c, k in zip(satir, ANAHTAR):
            if k in SAYISAL or k == "tevkifat":
                c.number_format = "#,##0.00"
            elif k in ("evrak_tarihi", "kayit_tarihi"):
                c.number_format = "dd\\.mm\\.yyyy"
    for i, w in enumerate([8, 15, 15, 12, 12, 8, 19, 13, 30, 22, 22, 40, 12, 12, 8, 15, 14, 12, 10, 30, 8, 12, 12, 9, 9, 12, 12, 10,
                           10, 10, 10, 40, 10], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    wb.save(yol)
    return yol
