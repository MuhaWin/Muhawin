"""KDV 1 / KDV 2 beyannamesi (GİB yeni beyanname görünümü) — kontrolü tamamlanan dönemin muavininden hazırlanır.

Kaynak: dönem klasöründeki rapor (kontrol sonuçları) ve muavin. Mükellef bilgileri firma kartından (firma.json),
düzenleyen / onaylayan bilgileri büro ayarlarından gelir. Hesaplama `cikti.kdv_beyani` / `cikti.kdv2_beyani` ile aynıdır;
burada yalnız beyanname kutularına yerleştirilir (JSON) — ekran ve PDF çıktı MuhaWin içindeki sayfada çizilir.
"""
from __future__ import annotations

import json
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from . import ayarlar, depo
from .cikti import AYLAR, kdv2_beyani, kdv_beyani
from .eslestirme import BILGI_DURUMLARI

SIFIR = Decimal("0")


def _f(x) -> float:
    return float(x or 0)


def _donem(v: dict) -> tuple[date, date]:
    d = v.get("donem", {})
    bas_s, bit_s = d.get("Alış") or d.get("Satış") or [None, None]
    bas = datetime.strptime(bas_s, "%Y-%m-%d").date() if bas_s else date.today().replace(day=1)
    bit = datetime.strptime(bit_s, "%Y-%m-%d").date() if bit_s else date.today()
    return bas, bit


def uygunluk(v: dict) -> dict:
    """Kontrol sonucu: sorunlu kayıt sayısı ve kullanıcı kararıyla uygun sayılanlar."""
    sorunlu = sum(d["adet"] for d in v.get("durumlar", []) if d.get("sorunlu"))
    karar = sum(d["adet"] for d in v.get("durumlar", []) if d.get("ad") == "Uygun (kullanıcı kararı)")
    fy = sum(1 for x in v.get("faturasi_olmayanlar", []) if x.get("uygun") is not None)
    return {"sorunlu": sorunlu, "kullanici_karari": karar + fy, "temiz": sorunlu == 0,
            "fatura_sayisi": v.get("fatura_sayisi", 0)}


def mukellef(firma_klasor: Path, v: dict) -> dict:
    bilgi = {}
    try:
        bilgi = json.loads((firma_klasor / "firma.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        pass
    buro = ayarlar.oku()
    return {"vkn": v.get("vkn") or bilgi.get("vkn") or "", "unvan": v.get("firma") or bilgi.get("ad") or firma_klasor.name,
            "eposta": bilgi.get("eposta", ""), "telefon": bilgi.get("telefon", ""),
            "vergi_dairesi": bilgi.get("vergi_dairesi", ""),
            "duzenleyen": {"ad": buro.get("ad_soyad", ""), "vkn": buro.get("vkn", ""), "eposta": buro.get("eposta", ""),
                           "telefon": buro.get("telefon", ""), "unvan": buro.get("buro_unvan", "")}}


def _satir(x) -> dict:
    return {"ad": x.ad, "matrah": None if x.matrah is None else _f(x.matrah), "oran": x.oran,
            "vergi": None if x.vergi is None else _f(x.vergi), "not": x.not_}


def kdv1_json(b, bas: date, bit: date) -> dict:
    devreden_tevkifat = SIFIR
    return {
        "donem": {"yil": bas.year, "ay": bas.month, "ay_adi": AYLAR[bas.month], "tip": "AYLIK",
                  "bas": bas.isoformat(), "bit": bit.isoformat()},
        "matrah": {"toplam_matrah": _f(b.toplam_matrah), "hesaplanan_kdv": _f(b.hesaplanan),
                   "ilave_kdv": _f(b.ilave_toplam), "toplam_kdv": _f(b.toplam_hesaplanan)},
        "indirimler": {"toplam": _f(b.indirimler_toplami), "onceki_donemden_devreden": _f(b.devreden),
                       "bu_donem": _f(b.bu_donem_indirim), "satirlar": [_satir(x) for x in b.indirimler],
                       "kdv2_odenen": _f(b.kdv2)},
        "istisnalar": {"teslim_toplami": _f(b.istisna_bedeli), "iade_edilebilir": _f(b.iade_edilebilir or 0),
                       "satirlar": [_satir(x) for x in b.istisnalar], "efatura": b.efatura_istisna},
        "sonuc": {"tecil": 0.0, "odenecek": _f(b.odenecek), "iade_edilecek": _f(b.iade_edilmesi_gereken),
                  "sonraki_doneme_devreden": _f(b.sonraki_doneme)},
        "diger": {"teslim_bedeli_aylik": _f(b.aylik_bedel), "teslim_bedeli_kumulatif": _f(b.kumulatif_bedel),
                  "kredi_karti": 0.0},
        "teslimler": [_satir(x) for x in b.teslimler],
        "kismi_tevkifat": [_satir(x) for x in b.kismi_tevkifat],
        "ilave": [_satir(x) for x in b.ilave],
        "satis_hesaplari": [_satir(x) for x in b.satis_hesaplari],
        "tahakkuk": {k: (_f(val) if isinstance(val, Decimal) else val) for k, val in (b.tahakkuk or {}).items()} if b.tahakkuk else None,
        "iade_cevap": b.iade_cevap, "iade_onerisi": None if b.iade_onerisi is None else _f(b.iade_onerisi),
        "notlar": b.notlar,
    }


def kdv2_json(k, bas: date, bit: date) -> dict:
    return {
        "donem": {"yil": bas.year, "ay": bas.month, "ay_adi": AYLAR[bas.month].upper(), "tip": "AYLIK"},
        "vergi_bildirimi": {"matrah": _f(k.matrah), "tevkif_edilen": _f(k.tevkif), "ilave": 0.0, "odenecek": _f(k.tevkif)},
        "kesintiler": [{"islem": x.islem, "odeme_turu": "KDV Genel Uygulama Tebliği Gereği Yapılan Ödemeler",
                        "unvan": x.unvan, "vkn": x.vkn, "belge_no": x.belge_no, "tarih": x.tarih.isoformat() if x.tarih else "",
                        "matrah": _f(x.matrah), "kdv_orani": x.kdv_orani, "tevkifat_orani": x.tevkifat_orani,
                        "tevkif": _f(x.tutar), "kaynak": x.kaynak} for x in k.satirlar],
        "ozet": [{"islem": a, "kdv_orani": o, "tevkifat_orani": t, "matrah": _f(m), "vergi": _f(v)} for a, o, t, m, v in k.ozet()],
        "muavinde_yok": k.muavinde_yok, "onceki_odemeler": _f(k.odemeler),
    }


def hazirla(rapor_id: str, iade: dict | None = None) -> dict:
    """Dönem raporu + muavin → {mukellef, uygunluk, kdv1, kdv2}."""
    from .kontrol import muavin_dosyalari, muavinleri_oku
    from .web.app import rapor_oku
    v = rapor_oku(rapor_id)
    k = depo.rapor_klasoru(rapor_id)
    if v is None or not k:
        raise FileNotFoundError("Rapor bulunamadı")
    if not muavin_dosyalari(k):
        raise FileNotFoundError("Bu dönemin muavini yok; önce muavini indirin ya da yükleyin")
    muavin, _ = muavinleri_oku(muavin_dosyalari(k))
    bas, bit = _donem(v)
    iade_yolu = k / "iade_talebi.json"
    if iade and iade.get("cevap") in ("evet", "hayir"):
        iade_yolu.write_text(json.dumps(iade, ensure_ascii=False), encoding="utf-8")
    else:
        try:
            iade = json.loads(iade_yolu.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            iade = None
    # KDV2: önceki aylarda kesilip bu ay kaydedilen tevkifatlı faturalar
    tevk = list(v.get("tevkifatlar", []))
    firma_ad, yil, ay = rapor_id.split("/")
    for geri in (1, 2):
        y, a = int(yil), int(ay) - geri
        if a < 1:
            y, a = y - 1, a + 12
        onceki = depo.rapor_klasoru(f"{firma_ad}/{y}/{a:02d}")
        if onceki and (onceki / depo.RAPOR_JSON).exists():
            try:
                tevk += json.loads((onceki / depo.RAPOR_JSON).read_text(encoding="utf-8")).get("tevkifatlar", [])
            except ValueError:
                pass
    b = kdv_beyani(muavin, bas, bit, v.get("faturalar"), iade)
    k2 = kdv2_beyani(muavin, bas, bit, tevk)
    out = {"id": rapor_id, "mukellef": mukellef(k.parent.parent, v), "uygunluk": uygunluk(v),
           "kdv1": kdv1_json(b, bas, bit), "kdv2": kdv2_json(k2, bas, bit),
           "hazirlandi": datetime.now().isoformat(timespec="seconds"), "kaynak": "kontrol"}
    try:
        (k / "beyanname.json").write_text(json.dumps(out, ensure_ascii=False), encoding="utf-8")
    except OSError:
        pass
    return out


def dosyadan(muavin_yollari: list[Path], program: str | None, bas: date, bit: date, firma: str, vkn: str) -> dict:
    """Muhasebe programı olmayan firmalar: yüklenen muavinden (tüm hesaplar) beyanname."""
    from .kontrol import muavinleri_oku
    muavin, bilgi = muavinleri_oku(muavin_yollari, program)
    b = kdv_beyani(muavin, bas, bit, None, None)
    k2 = kdv2_beyani(muavin, bas, bit, None)
    buro = ayarlar.oku()
    return {"id": None, "mukellef": {"vkn": vkn, "unvan": firma or bilgi.get("firma") or "", "eposta": "", "telefon": "", "vergi_dairesi": "",
                                     "duzenleyen": {"ad": buro.get("ad_soyad", ""), "vkn": buro.get("vkn", ""), "eposta": buro.get("eposta", ""),
                                                    "telefon": buro.get("telefon", ""), "unvan": buro.get("buro_unvan", "")}},
            "uygunluk": {"sorunlu": 0, "kullanici_karari": 0, "temiz": True, "fatura_sayisi": 0, "kontrolsuz": True},
            "kdv1": kdv1_json(b, bas, bit), "kdv2": kdv2_json(k2, bas, bit),
            "hazirlandi": datetime.now().isoformat(timespec="seconds"), "kaynak": "dosya", "muavin_satir": len(muavin)}


def excel(out: dict, tur: str) -> bytes:
    """Beyanname kutularını Excel'e döker (tek sayfa; KDV 1 ya da KDV 2)."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    m = out.get("mukellef", {})
    kalin = Font(bold=True)
    bas_dolgu = PatternFill("solid", fgColor="E9EEF6")

    def baslik(t):
        ws.append([])
        ws.append([t])
        ws.cell(ws.max_row, 1).font = kalin
        ws.cell(ws.max_row, 1).fill = bas_dolgu

    def satirlar(rows, para_kol=()):
        for r in rows:
            ws.append(r)
            for c in para_kol:
                ws.cell(ws.max_row, c).number_format = "#,##0.00"
                ws.cell(ws.max_row, c).alignment = Alignment(horizontal="right")

    if tur == "kdv2":
        k = out["kdv2"]
        ws.title = "KDV 2"
        ws.append(["Katma Değer Vergisi Beyannamesi (Vergi Sorumluları İçin)"]); ws.cell(1, 1).font = Font(bold=True, size=13)
        ws.append([f"{m.get('unvan', '')} · VKN {m.get('vkn', '')} · {k['donem']['yil']} / {k['donem']['ay_adi']}"])
        baslik("VERGİ BİLDİRİMİ")
        v = k["vergi_bildirimi"]
        satirlar([["KDV Matrahı", v["matrah"]], ["Tevkif Edilen KDV", v["tevkif_edilen"]], ["İlave Edilecek KDV", v["ilave"]], ["Ödenmesi Gereken KDV", v["odenecek"]]], (2,))
        baslik("KDV Kesintisi Yapılanlara Ait Bildirim")
        ws.append(["İşlem Türü", "Ünvan", "TCKN / VKN", "Belge No", "Tarih", "Matrah", "KDV Oranı", "Tevkifat Oranı", "Tevkif Edilen KDV"])
        satirlar([[x["islem"], x["unvan"], x["vkn"], x["belge_no"], x["tarih"], x["matrah"], x["kdv_orani"], x["tevkifat_orani"], x["tevkif"]] for x in k["kesintiler"]], (6, 9))
        baslik("Kısmi Tevkifat Uygulanan İşlemlere Ait Bildirim")
        ws.append(["İşlem Türü", "Matrah", "KDV Oranı", "Tevkifat Oranı", "Vergi"])
        satirlar([[x["islem"], x["matrah"], x["kdv_orani"], x["tevkifat_orani"], x["vergi"]] for x in k["ozet"]], (2, 5))
    else:
        b = out["kdv1"]
        ws.title = "KDV 1"
        ws.append(["Katma Değer Vergisi Beyannamesi (Gerçek Usulde Vergilendirilen Mükellefler İçin)"]); ws.cell(1, 1).font = Font(bold=True, size=13)
        ws.append([f"{m.get('unvan', '')} · VKN {m.get('vkn', '')} · {b['donem']['yil']} / {b['donem']['ay_adi']}"])
        baslik("MATRAH")
        satirlar([["Toplam Matrah", b["matrah"]["toplam_matrah"]], ["Hesaplanan KDV", b["matrah"]["hesaplanan_kdv"]],
                  ["Daha Önce İndirim Konusu Yapılan KDV'nin İlavesi", b["matrah"]["ilave_kdv"]], ["Toplam KDV", b["matrah"]["toplam_kdv"]]], (2,))
        baslik("İNDİRİMLER")
        satirlar([["İndirimler Toplamı", b["indirimler"]["toplam"]], ["Önceki Dönemden Devreden KDV", b["indirimler"]["onceki_donemden_devreden"]],
                  ["Bu Döneme Ait İndirilecek KDV", b["indirimler"]["bu_donem"]], ["Sorumlu Sıfatıyla Ödenen KDV (109)", b["indirimler"]["kdv2_odenen"]]], (2,))
        baslik("İSTİSNALAR")
        satirlar([["İstisna Kapsamına Giren Toplam Teslim", b["istisnalar"]["teslim_toplami"]], ["İade Edilebilir KDV", b["istisnalar"]["iade_edilebilir"]]], (2,))
        baslik("SONUÇ HESAPLARI")
        satirlar([["Tecil Edilecek KDV", b["sonuc"]["tecil"]], ["Ödenmesi Gereken KDV", b["sonuc"]["odenecek"]],
                  ["İade Edilmesi Gereken KDV", b["sonuc"]["iade_edilecek"]], ["Sonraki Döneme Devreden KDV", b["sonuc"]["sonraki_doneme_devreden"]]], (2,))
        baslik("DİĞER BİLGİLER")
        satirlar([["Teslim ve Hizmet Bedeli (Aylık)", b["diger"]["teslim_bedeli_aylik"]], ["Teslim ve Hizmet Bedeli (Kümülatif)", b["diger"]["teslim_bedeli_kumulatif"]],
                  ["Kredi Kartı ile Tahsil Edilen", b["diger"]["kredi_karti"]]], (2,))
        baslik("MATRAH DETAYI — Tevkifat Uygulanmayan İşlemler")
        ws.append(["İşlem Türü", "Matrah", "KDV Oranı", "Vergi"])
        satirlar([[x["ad"], x["matrah"], x["oran"], x["vergi"]] for x in b["teslimler"]], (2, 4))
        if b["kismi_tevkifat"]:
            baslik("MATRAH DETAYI — Kısmi Tevkifat Uygulanan İşlemler")
            ws.append(["İşlem Türü", "Matrah", "KDV Oranı", "Vergi"])
            satirlar([[x["ad"], x["matrah"], x["oran"], x["vergi"]] for x in b["kismi_tevkifat"]], (2, 4))
        baslik("İNDİRİMLER DETAYI — Oranlara Göre Dağılım")
        ws.append(["İndirim Türü", "KDV Oranı", "Alınan Mal ve Hizmet Bedeli", "KDV Tutarı", "Hesap"])
        satirlar([[x["ad"], x["oran"], x["matrah"], x["vergi"], x["not"]] for x in b["indirimler"]["satirlar"]], (3, 4))
        if b["istisnalar"]["satirlar"]:
            baslik("İSTİSNALAR DETAYI — Tam İstisna")
            ws.append(["İstisna Türü", "Teslim ve Hizmet Tutarı", "Yüklenilen KDV"])
            satirlar([[x["ad"], x["matrah"], b["istisnalar"]["iade_edilebilir"]] for x in b["istisnalar"]["satirlar"]], (2, 3))
        if b.get("notlar"):
            baslik("NOTLAR")
            for n in b["notlar"]:
                ws.append([n])
    for col, w in zip("ABCDEFGHI", (52, 34, 18, 18, 14, 16, 12, 14, 18)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- Toplu beyanname (tüm firmalar, bir dönem)
def toplu(yil: int, ay: int, yalniz_temiz: bool = False) -> dict:
    """Bir dönem için kontrolü yapılmış tüm firmaların KDV 1 / KDV 2 kutuları tek listede.
    Beyanname daha önce hazırlandıysa beyanname.json'dan okunur, yoksa muavinden hazırlanır; muavini olmayan ya da
    hazırlanamayan firmalar 'hata' alanıyla listelenir (gönderim öncesi kontrol listesi)."""
    from .kontrol import muavin_dosyalari
    from .web.app import _db
    satirlar = []
    with _db() as con:
        rows = con.execute("SELECT id, firma, vkn, bas, bit, fatura_sayisi, sorunlu FROM calisma ORDER BY firma").fetchall()
    gor = set()
    for rid, firma, vkn, bas, bit, fs, sorunlu in rows:
        if rid in gor or not str(bas or "").startswith(f"{yil:04d}-{ay:02d}"):
            continue
        gor.add(rid)
        if yalniz_temiz and sorunlu:
            continue
        s = {"id": rid, "firma": firma, "vkn": vkn, "sorunlu": int(sorunlu or 0), "fatura_sayisi": fs}
        k = depo.rapor_klasoru(rid)
        if not k or not muavin_dosyalari(k):
            s["hata"] = "muavin yok"
            satirlar.append(s)
            continue
        try:
            try:
                out = json.loads((k / "beyanname.json").read_text(encoding="utf-8"))
            except (OSError, ValueError):
                out = hazirla(rid)
            b, k2 = out["kdv1"], out["kdv2"]
            s.update({
                "hazirlandi": out.get("hazirlandi"),
                "toplam_matrah": b["matrah"]["toplam_matrah"], "hesaplanan_kdv": b["matrah"]["toplam_kdv"],
                "onceki_devreden": b["indirimler"]["onceki_donemden_devreden"], "bu_donem_indirim": b["indirimler"]["bu_donem"],
                "indirimler_toplami": b["indirimler"]["toplam"], "istisna": b["istisnalar"]["teslim_toplami"],
                "odenecek": b["sonuc"]["odenecek"], "iade": b["sonuc"]["iade_edilecek"], "devreden": b["sonuc"]["sonraki_doneme_devreden"],
                "bedel_aylik": b["diger"]["teslim_bedeli_aylik"], "bedel_kumulatif": b["diger"]["teslim_bedeli_kumulatif"],
                "kdv2_matrah": k2["vergi_bildirimi"]["matrah"], "kdv2_tevkif": k2["vergi_bildirimi"]["tevkif_edilen"],
                "iade_cevap_bekliyor": bool(((b.get("istisnalar") or {}).get("satirlar")) and not b.get("iade_cevap")),
            })
        except Exception as e:  # noqa: BLE001 — bir firmanın hatası diğerlerini durdurmasın
            s["hata"] = str(e)[:200]
        satirlar.append(s)
    top = {k: round(sum(float(s.get(k) or 0) for s in satirlar if "hata" not in s), 2)
           for k in ("toplam_matrah", "hesaplanan_kdv", "indirimler_toplami", "odenecek", "devreden", "kdv2_tevkif")}
    return {"yil": yil, "ay": ay, "ay_adi": AYLAR[ay], "satirlar": satirlar, "toplam": top,
            "hazirlandi": datetime.now().isoformat(timespec="seconds")}


TOPLU_KOLONLAR = [("firma", "Firma"), ("vkn", "VKN"), ("sorunlu", "Sorunlu kayıt"), ("toplam_matrah", "Toplam Matrah"),
                  ("hesaplanan_kdv", "Hesaplanan KDV"), ("onceki_devreden", "Önceki Dönemden Devreden"),
                  ("bu_donem_indirim", "Bu Dönem İndirilecek"), ("indirimler_toplami", "İndirimler Toplamı"),
                  ("istisna", "İstisna Teslim"), ("odenecek", "Ödenecek KDV"), ("iade", "İade Edilecek"),
                  ("devreden", "Sonraki Döneme Devreden"), ("bedel_aylik", "Bedel (Aylık)"), ("bedel_kumulatif", "Bedel (Kümülatif)"),
                  ("kdv2_matrah", "KDV2 Matrah"), ("kdv2_tevkif", "KDV2 Tevkif Edilen"), ("hata", "Not")]


def toplu_excel(t: dict) -> bytes:
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "KDV Özet"
    ws.append([f"KDV Beyanname Özeti · {t['yil']} / {t['ay_adi']} · {len(t['satirlar'])} firma · hazırlandı {t['hazirlandi'][:16].replace('T', ' ')}"])
    ws.cell(1, 1).font = Font(bold=True, size=13)
    ws.append([ad for _, ad in TOPLU_KOLONLAR])
    for c in range(1, len(TOPLU_KOLONLAR) + 1):
        ws.cell(2, c).font = Font(bold=True)
        ws.cell(2, c).fill = PatternFill("solid", fgColor="E9EEF6")
    para = {i + 1 for i, (k, _) in enumerate(TOPLU_KOLONLAR) if k not in ("firma", "vkn", "sorunlu", "hata")}
    for s in t["satirlar"]:
        ws.append([s.get(k, "") if s.get(k) is not None else "" for k, _ in TOPLU_KOLONLAR])
        for c in para:
            ws.cell(ws.max_row, c).number_format = "#,##0.00"
            ws.cell(ws.max_row, c).alignment = Alignment(horizontal="right")
    ws.append([])
    ws.append(["TOPLAM", "", ""] + [t["toplam"].get(k, "") if k in t["toplam"] else "" for k, _ in TOPLU_KOLONLAR[3:]])
    for c in range(1, len(TOPLU_KOLONLAR) + 1):
        ws.cell(ws.max_row, c).font = Font(bold=True)
        if c in para:
            ws.cell(ws.max_row, c).number_format = "#,##0.00"
    ws.column_dimensions["A"].width = 38
    for col in "BCDEFGHIJKLMNOPQ":
        ws.column_dimensions[col].width = 18
    ws.freeze_panes = "B3"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- GİB beyanname PDF'i ile karşılaştırma
# GİB e-Beyanname çıktısındaki (KDV1 / KDV2) kutu adları → MuhaWin beyanname alanı. Metin, MuhaWin'de pdf.js ile çıkarılıp
# gönderilir (Python tarafında PDF kütüphanesi gerekmez). Etiketten sonraki ilk Türkçe biçimli sayı alınır.
_GIB_KUTULAR = [
    # (anahtar, görünen ad, [etiket düzenli ifadeleri], tür, yol)
    ("toplam_matrah", "Toplam Matrah", [r"Toplam\s+Matrah"], "kdv1", ("matrah", "toplam_matrah")),
    ("hesaplanan_kdv", "Hesaplanan KDV", [r"Hesaplanan\s+(?:Katma\s+De[ğg]er\s+Vergisi|KDV)(?!\s*n)"], "kdv1", ("matrah", "hesaplanan_kdv")),
    ("ilave_kdv", "İlave Edilecek KDV", [r"[İI]lavesi", r"[İI]lave\s+Edilecek\s+KDV"], "kdv1", ("matrah", "ilave_kdv")),
    ("toplam_kdv", "Toplam KDV", [r"Toplam\s+(?:Katma\s+De[ğg]er\s+Vergisi|KDV)(?!\s+Matrah)"], "kdv1", ("matrah", "toplam_kdv")),
    ("indirimler_toplami", "İndirimler Toplamı", [r"[İI]ndirimler\s+Toplam[ıi]"], "kdv1", ("indirimler", "toplam")),
    ("onceki_devreden", "Önceki Dönemden Devreden KDV", [r"[ÖO]nceki\s+D[öo]nemden\s+Devreden"], "kdv1", ("indirimler", "onceki_donemden_devreden")),
    ("bu_donem_indirim", "Bu Döneme Ait İndirilecek KDV", [r"Bu\s+D[öo]neme\s+Ait\s+[İI]ndirilecek\s+KDV"], "kdv1", ("indirimler", "bu_donem")),
    ("istisna", "İstisna Teslim Toplamı", [r"[İI]stisna\s+Kapsam[ıi]na\s+Giren[^0-9]{0,80}Tutar[ıi]"], "kdv1", ("istisnalar", "teslim_toplami")),
    ("tecil", "Tecil Edilecek KDV", [r"Tecil\s+Edilecek"], "kdv1", ("sonuc", "tecil")),
    ("odenecek", "Ödenmesi Gereken KDV", [r"[ÖO]denmesi\s+Gereken\s+(?:Katma\s+De[ğg]er\s+Vergisi|KDV)"], "kdv1", ("sonuc", "odenecek")),
    ("iade", "İade Edilmesi Gereken KDV", [r"[İI]ade\s+Edilmesi\s+Gereken"], "kdv1", ("sonuc", "iade_edilecek")),
    ("devreden", "Sonraki Döneme Devreden KDV", [r"Sonraki\s+D[öo]neme\s+Devreden"], "kdv1", ("sonuc", "sonraki_doneme_devreden")),
    ("bedel_aylik", "Teslim Bedeli (Aylık)", [r"Bedel\s*\(?\s*Ayl[ıi]k"], "kdv1", ("diger", "teslim_bedeli_aylik")),
    ("bedel_kumulatif", "Teslim Bedeli (Kümülatif)", [r"Bedel\s*\(?\s*K[üu]m[üu]latif"], "kdv1", ("diger", "teslim_bedeli_kumulatif")),
    ("kredi_karti", "Kredi Kartı ile Tahsil", [r"Kredi\s+Kart[ıi]\s+[İI]le\s+Tahsil[^0-9]{0,120}Bedel"], "kdv1", ("diger", "kredi_karti")),
    ("kdv2_matrah", "KDV2 — KDV Matrahı", [r"KDV\s+Matrah[ıi]"], "kdv2", ("vergi_bildirimi", "matrah")),
    ("kdv2_tevkif", "KDV2 — Tevkif Edilen KDV", [r"Tevkif\s+Edilen\s+KDV"], "kdv2", ("vergi_bildirimi", "tevkif_edilen")),
    ("kdv2_odenecek", "KDV2 — Ödenmesi Gereken KDV", [r"[ÖO]denmesi\s+Gereken\s+(?:Katma\s+De[ğg]er\s+Vergisi|KDV)"], "kdv2", ("vergi_bildirimi", "odenecek")),
]
_SAYI = re.compile(r"(?<![\d.,])(-?\d{1,3}(?:\.\d{3})*(?:,\d{1,2})?|-?\d+,\d{1,2})(?![\d.])")


def gib_sayi(s: str) -> float | None:
    s = s.strip()
    if not s:
        return None
    try:
        return float(s.replace(".", "").replace(",", "."))
    except ValueError:
        return None


def gib_metin_coz(metin: str, tur: str) -> dict:
    """GİB beyanname PDF metninden kutu değerleri: {anahtar: tutar}. Bulunamayanlar yoktur."""
    m = re.sub(r"[ \t ]+", " ", metin or "")
    out = {}
    for anahtar, _ad, etiketler, t, _yol in _GIB_KUTULAR:
        if t != tur:
            continue
        for et in etiketler:
            for mm in re.finditer(et, m, flags=re.I):
                kuyruk = m[mm.end(): mm.end() + 160]
                sm = _SAYI.search(kuyruk)
                if sm and gib_sayi(sm.group(1)) is not None:
                    out[anahtar] = gib_sayi(sm.group(1))
                    break
            if anahtar in out:
                break
    # dönem / VKN
    dm = re.search(r"Y[ıi]l\s*:?\s*(\d{4}).{0,40}?Ay\s*:?\s*([A-ZÇĞİÖŞÜa-zçğıöşü]+)", m)
    if dm:
        out["_donem"] = f"{dm.group(1)} {dm.group(2)}"
    vm = re.search(r"Vergi\s+Kimlik\s+No\s*:?\s*(\d{10,11})", m)
    if vm:
        out["_vkn"] = vm.group(1)
    return out


def karsilastir(rapor_id: str, metin: str, tur: str = "kdv1") -> dict:
    """MuhaWin beyannamesi (kontrol edilen muavinden) ile GİB'den indirilen beyanname PDF metnini kutu kutu karşılaştırır."""
    k = depo.rapor_klasoru(rapor_id)
    if not k:
        raise FileNotFoundError("Rapor bulunamadı")
    try:
        out = json.loads((k / "beyanname.json").read_text(encoding="utf-8"))
    except (OSError, ValueError):
        out = hazirla(rapor_id)
    tur = "kdv2" if tur == "kdv2" else "kdv1"
    gib = gib_metin_coz(metin, tur)
    sat = []
    for anahtar, ad, _et, t, yol in _GIB_KUTULAR:
        if t != tur:
            continue
        bizim = out[tur]
        for p in yol:
            bizim = (bizim or {}).get(p) if isinstance(bizim, dict) else None
        bizim = None if bizim is None else round(float(bizim), 2)
        g = gib.get(anahtar)
        fark = None if (g is None or bizim is None) else round(bizim - g, 2)
        durum = "bulunamadı" if g is None else ("uyumlu" if abs(fark or 0) < 0.01 else "farklı")
        sat.append({"anahtar": anahtar, "ad": ad, "muhawin": bizim, "gib": g, "fark": fark, "durum": durum})
    farkli = [s for s in sat if s["durum"] == "farklı"]
    bulunan = [s for s in sat if s["gib"] is not None]
    mk = out.get("mukellef", {})
    uyari = []
    if gib.get("_vkn") and mk.get("vkn") and gib["_vkn"] != str(mk["vkn"]):
        uyari.append(f"PDF'teki VKN ({gib['_vkn']}) bu firmanın VKN'si ({mk['vkn']}) ile aynı değil.")
    if not bulunan:
        uyari.append("PDF'te beyanname kutuları bulunamadı; metin katmanı olmayan (taranmış) bir PDF olabilir.")
    return {"id": rapor_id, "tur": tur, "satirlar": sat, "farkli": len(farkli), "bulunan": len(bulunan),
            "gib_donem": gib.get("_donem"), "gib_vkn": gib.get("_vkn"), "uyarilar": uyari,
            "sonuc": "uyumlu" if bulunan and not farkli else ("farklı" if farkli else "okunamadı")}
