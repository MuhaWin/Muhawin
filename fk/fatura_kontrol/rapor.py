"""Excel rapor üretimi."""
from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .eslestirme import (fatura_tipi, TUM_DURUMLAR, UYUMLU, DONEM_FARKI, DONEM_DISI, KDVSIZ, KURUS_FARKI, Rapor)

TL = '#,##0.00;[Red]-#,##0.00'
BASLIK = PatternFill("solid", fgColor="1F3864")
RENK = {
    UYUMLU: "E2EFDA",
    DONEM_FARKI: "FFF2CC",
    DONEM_DISI: "EDEDED",
    KDVSIZ: "EDEDED",
    KURUS_FARKI: "FFF2CC",
    "Red / iptal – muavinde kaydı yok, uyumlu": "E2EFDA",
}
SORUNLU = "F8CBAD"


def _tablo(ws, basliklar, satirlar, para_kolonlari=(), genislik=None):
    ws.append(basliklar)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = BASLIK
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in satirlar:
        ws.append(r)
    for col in para_kolonlari:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = TL
    ws.freeze_panes = "A2"
    if satirlar:
        ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(genislik or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def excel_yaz(rapor: Rapor, faturalar, firma: str, vkn: str, uyarilar: list[str], yol: Path) -> Path:
    wb = Workbook()

    # --- Özet
    ws = wb.active
    ws.title = "Özet"
    ws["A1"] = "Fatura – Muavin Kontrol Raporu"
    ws["A1"].font = Font(bold=True, size=14)
    bilgiler = [
        ("Firma", firma), ("VKN", vkn),
        ("Rapor tarihi", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Okunan fatura", len(faturalar)),
    ]
    for yon, (b, s) in rapor.donem.items():
        bilgiler.append((f"Kontrol aralığı ({'Alış' if yon == 'ALIS' else 'Satış'})",
                         f"{b:%d.%m.%Y} – {s:%d.%m.%Y}"))
    bilgiler.append(("Yapılan kontroller", "; ".join(rapor.calisan_kontroller) or "—"))
    for i, (a, d) in enumerate(bilgiler, start=3):
        ws.cell(i, 1, a).font = Font(bold=True)
        ws.cell(i, 2, d)
    r = len(bilgiler) + 4
    ws.cell(r, 1, "Durum").font = Font(bold=True)
    ws.cell(r, 2, "Adet").font = Font(bold=True)
    sayac = Counter(s.durum for s in rapor.sonuclar)
    sayac["Faturası bulunamadı"] = sum(1 for s in rapor.faturasi_olmayanlar if getattr(s, "uygun", None) is None)
    for d in TUM_DURUMLAR:
        r += 1
        ws.cell(r, 1, d)
        c = ws.cell(r, 2, sayac.get(d, 0))
        if sayac.get(d, 0):
            ws.cell(r, 1).fill = PatternFill("solid", fgColor=RENK.get(d, SORUNLU))
    r += 1
    ws.cell(r, 1, "Fatura dışı kayıt (bilgi)")
    ws.cell(r, 2, len(rapor.fatura_disi))
    if uyarilar:
        r += 2
        ws.cell(r, 1, "Uyarılar").font = Font(bold=True)
        for u in uyarilar:
            r += 1
            ws.cell(r, 1, u)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 70

    # --- Fatura bazlı kontrol
    ws = wb.create_sheet("Fatura Kontrol")
    sirala = {d: i for i, d in enumerate(TUM_DURUMLAR)}
    sonuclar = sorted(rapor.sonuclar, key=lambda s: (-sirala[s.durum], s.fatura.tarih, s.fatura.belge_no))
    satirlar = []
    for s in sonuclar:
        f = s.fatura
        satirlar.append([
            s.durum, "Alış" if f.yon == "ALIS" else "Satış", s.kontrol, f.belge_no,
            f.tarih, f.karsi_unvan, f.karsi_vkn,
            ", ".join(f"%{o}" for o in sorted(f.kdv_oranlari)),
            float(s.fatura_tutari), float(s.muavin_tutari), float(s.fark),
            ", ".join(sorted({x.hesap_kodu for x in s.satirlar})),
            ", ".join(sorted({x.fis_no for x in s.satirlar})),
            ", ".join(sorted({x.fis_tarihi.strftime('%d.%m.%Y') for x in s.satirlar})),
            s.aciklama, fatura_tipi(f), f.profil, f.tip,
        ])
    _tablo(ws, ["Durum", "Yön", "Kontrol", "Belge No", "Fatura Tarihi", "Karşı Taraf", "VKN/TCKN",
                "KDV Oranları", "Fatura Tutarı", "Muavin Tutarı", "Fark", "Muavin Hesap",
                "Fiş No", "Fiş Tarihi", "Açıklama", "Fatura Tipi / İstisna", "Senaryo", "Tip"],
           satirlar, para_kolonlari=(9, 10, 11),
           genislik=[30, 7, 9, 19, 12, 42, 13, 11, 14, 14, 12, 22, 9, 12, 50, 36, 14, 12])
    for row in ws.iter_rows(min_row=2):
        row[0].fill = PatternFill("solid", fgColor=RENK.get(row[0].value, SORUNLU))
        row[4].number_format = "DD.MM.YYYY"

    # --- Muavinde olup faturası olmayan
    ws = wb.create_sheet("Faturası Bulunamayan")
    _tablo(ws, ["Hesap", "Hesap Adı", "Fiş Tarihi", "Fiş No", "Belge Tarihi", "Belge No",
                "Ünvan", "Tutar", "Açıklama", "Kullanıcı Kararı"],
           [[s.hesap_kodu, s.hesap_adi, s.fis_tarihi, s.fis_no, s.belge_tarihi, s.belge_no,
             s.unvan, float(s.net), s.aciklama, ("Uygun: " + s.uygun) if getattr(s, "uygun", None) is not None else ""]
            for s in sorted(rapor.faturasi_olmayanlar, key=lambda s: (s.belge_tarihi or s.fis_tarihi, s.belge_no))],
           para_kolonlari=(8,), genislik=[12, 28, 12, 9, 12, 19, 45, 14, 70, 50])
    for row in ws.iter_rows(min_row=2):
        row[2].number_format = row[4].number_format = "DD.MM.YYYY"

    # --- Fatura dışı kayıtlar
    ws = wb.create_sheet("Fatura Dışı Kayıtlar")
    _tablo(ws, ["Hesap", "Hesap Adı", "Fiş Tarihi", "Fiş No", "Tutar", "Açıklama"],
           [[s.hesap_kodu, s.hesap_adi, s.fis_tarihi, s.fis_no, float(s.net), s.aciklama]
            for s in rapor.fatura_disi],
           para_kolonlari=(5,), genislik=[12, 28, 12, 9, 14, 80])
    for row in ws.iter_rows(min_row=2):
        row[2].number_format = "DD.MM.YYYY"

    # --- Red / iptal edilmiş faturalar (Mikro listeleri)
    red = getattr(rapor, "red_iptal", None) or []
    if red:
        ws = wb.create_sheet("Red - İptal")
        _tablo(ws, ["Sonuç", "Red / İptal", "Belge No", "Tarih", "Karşı Taraf", "VKN/TCKN", "Matrah", "KDV", "Yekün", "Döviz",
                    "Fatura Türü", "Durum (liste)", "e-Fatura Dosyalarında", "Muavinde", "Hesap", "Fiş", "Liste Dosyası"],
               [[b["sonuc"], b.get("ne", ""), b["belge_no"], b.get("tarih"), b.get("unvan", ""), b.get("vkn", ""),
                 float(b.get("matrah") or 0), float(b.get("kdv") or 0), float(b.get("yekun") or 0), b.get("doviz", ""),
                 b.get("tur", ""), b.get("durum", ""), "Var" if b["faturada"] else "Yok", "Kayıtlı" if b["muavinde"] else "Kayıt yok",
                 ", ".join(sorted({x.hesap_kodu for x in b["satirlar"]})),
                 ", ".join(sorted({f"{x.fis_no} ({x.fis_tarihi:%d.%m.%Y})" for x in b["satirlar"]})), b.get("dosya", "")]
                for b in red],
               para_kolonlari=(7, 8, 9), genislik=[40, 11, 19, 12, 45, 13, 14, 12, 14, 7, 12, 30, 12, 11, 18, 26, 30])
        for row in ws.iter_rows(min_row=2):
            row[3].number_format = "DD.MM.YYYY"
            row[0].fill = PatternFill("solid", fgColor=RENK.get(row[0].value, SORUNLU))

    # --- Okunan faturalar
    ws = wb.create_sheet("Okunan Faturalar")
    _tablo(ws, ["Yön", "Belge No", "Tarih", "Senaryo", "Tip", "Karşı Taraf", "VKN/TCKN",
                "Matrah", "KDV", "Tevkifat", "Vergiler Dahil", "Ödenecek", "KDV Dökümü",
                "Diğer Vergiler", "ETTN", "Dosya"],
           [["Alış" if f.yon == "ALIS" else "Satış", f.belge_no, f.tarih, f.profil, f.tip,
             f.karsi_unvan, f.karsi_vkn, float(f.matrah), float(f.kdv_toplam), float(f.tevkifat),
             float(f.vergiler_dahil), float(f.odenecek),
             "; ".join(f"%{o}: {v}" for o, v in sorted(f.kdv_oranlari.items())),
             "; ".join(f"{a}: {v}" for a, v in f.diger_vergiler.items()), f.ettn, f.kaynak]
            for f in sorted(faturalar, key=lambda f: (f.tarih, f.belge_no))],
           para_kolonlari=(8, 9, 10, 11, 12),
           genislik=[7, 19, 12, 14, 10, 45, 13, 14, 12, 11, 14, 14, 24, 28, 38, 50])
    for row in ws.iter_rows(min_row=2):
        row[2].number_format = "DD.MM.YYYY"

    # --- KDV tevkifatı (KDV2)
    tev = tevkifat_satirlari(faturalar)
    if tev:
        ws = wb.create_sheet("KDV Tevkifatı", 1)
        _tablo(ws, ["Yön", "Tarih", "Belge No", "Karşı Taraf", "VKN/TCKN", "Matrah", "KDV", "Tevkifat Kodu",
                    "Tevkifat Türü", "Oran", "Tevkif Edilen KDV (KDV2)", "Ödenecek"],
               [[t["yon"], t["tarih"], t["belge_no"], t["karsi_taraf"], t["vkn"], t["matrah"], t["kdv"],
                 t["kod"], t["tur"], t["oran"], t["tutar"], t["odenecek"]] for t in tev],
               para_kolonlari=(6, 7, 11, 12), genislik=[7, 12, 19, 45, 13, 14, 13, 10, 40, 7, 16, 14])
        for row in ws.iter_rows(min_row=2):
            row[1].number_format = "DD.MM.YYYY"
        for yon in ("Alış", "Satış"):
            t = sum(x["tutar"] for x in tev if x["yon"] == yon)
            if t:
                ws.append([])
                ws.append([f"{yon} toplamı (KDV2)" if yon == "Alış" else "Satış toplamı (müşterinin tevkif ettiği)",
                           None, None, None, None, None, None, None, None, None, t])
                ws.cell(ws.max_row, 1).font = Font(bold=True)
                ws.cell(ws.max_row, 11).font = Font(bold=True)
                ws.cell(ws.max_row, 11).number_format = "#,##0.00"

    wb.save(yol)
    return yol


def doviz_satirlari(faturalar, rapor=None) -> list[dict]:
    """Dövizli faturalar: döviz tutarı, kur ve TL karşılığı (karşılaştırma TL karşılığıyla yapılır)."""
    durum: dict[str, list[str]] = {}
    for s in (rapor.sonuclar if rapor else []):
        durum.setdefault(s.fatura.belge_no, []).append(f"{s.kontrol}: {s.durum}")
    return [{"yon": "Alış" if f.yon == "ALIS" else "Satış", "tarih": f.tarih, "belge_no": f.belge_no,
             "karsi_taraf": f.karsi_unvan, "vkn": f.karsi_vkn, "doviz": f.doviz,
             "kur": float(f.kur) if f.kur else None,
             "doviz_matrah": float(f.doviz_matrah or 0), "doviz_toplam": float(f.doviz_odenecek or 0),
             "matrah": float(f.matrah), "kdv": float(f.kdv_toplam), "toplam": float(f.odenecek),
             "durum": "; ".join(durum.get(f.belge_no, []))}
            for f in sorted(faturalar, key=lambda f: (f.yon, f.tarih, f.belge_no)) if f.doviz]


def tevkifat_satirlari(faturalar) -> list[dict]:
    """Tevkifatlı ALIŞ faturalarının her tevkifat kalemi için bir satır (KDV2: sorumlu sıfatıyla beyan; satış
    faturalarındaki tevkifatı alıcı beyan eder, bu yüzden listelenmez)."""
    satirlar = []
    for f in sorted((f for f in faturalar if f.yon == "ALIS"), key=lambda f: (f.tarih, f.belge_no)):
        detay = f.tevkifat_detay or ([{"kod": "", "ad": "", "oran": "", "kdv": f.kdv_toplam, "tutar": f.tevkifat}]
                                     if f.tevkifat else [])
        for d in detay:
            satirlar.append({"yon": "Alış" if f.yon == "ALIS" else "Satış", "tarih": f.tarih, "belge_no": f.belge_no,
                             "karsi_taraf": f.karsi_unvan, "vkn": f.karsi_vkn, "matrah": float(f.matrah),
                             "kdv": float(d["kdv"]), "kod": d["kod"], "tur": d["ad"], "oran": d["oran"],
                             "tutar": float(d["tutar"]), "odenecek": float(f.odenecek)})
    return satirlar
