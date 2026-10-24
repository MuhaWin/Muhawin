"""Toplu kontrol ekranının API'si: firma listesi, iş başlat/durdur, canlı ilerleme, aylık geçmiş.

Luca otomasyonu ayrı bir süreçte (python -m fatura_kontrol.luca.toplu --sessiz ...) çalışır;
ilerleme o sürecin yazdığı durum.json ve çıktı günlüğünden okunur. Aynı anda tek iş çalışır.
"""
from __future__ import annotations

import json
import os
import re
import subprocess
import sys
import threading
import time
from datetime import date, datetime
from pathlib import Path

from fastapi import APIRouter, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel

class _TopluTembel:
    """luca.toplu (Playwright zinciri) ilk kullanımda yüklenir: web.app Playwright olmadan da içe aktarılabilir
    (tarayıcı köprüsü / Pyodide); toplu uçları çağrılınca masaüstündeki gibi çalışır."""
    def __getattr__(self, ad):
        from ..luca import toplu
        return getattr(toplu, ad)


T = _TopluTembel()
from .. import depo

router = APIRouter(prefix="/api/toplu")
_kilit = threading.Lock()
_is: dict = {}          # {"tur", "proc", "klasor", "log", "basladi"}
# MuhaWin'deki Luca oturumu (çerezler): yalnız bellekte tutulur, hiçbir yanıtta geri verilmez; Luca otomasyonu başlarken
# süreç için geçici bir dosyaya (yalnız kullanıcı okuyabilir) yazılır, otomasyon okuyunca siler.
_LUCA_OTURUM: dict = {"cerez": [], "zaman": 0.0}
LUCA_OTURUM_OMRU = 6 * 3600
MODUL = "fatura_kontrol.luca.toplu"     # testlerde sahte modülle değiştirilir


def _toplu_kok() -> Path:
    return depo.toplu_kok()


def _calisiyor() -> bool:
    p = _is.get("proc")
    return bool(p and p.poll() is None)


def _baslat(tur: str, args: list[str], klasor: Path):
    with _kilit:
        if _calisiyor():
            raise HTTPException(409, "Zaten çalışan bir işlem var. Bitmesini bekleyin ya da durdurun.")
        klasor.mkdir(parents=True, exist_ok=True)
        (klasor / "durum.json").unlink(missing_ok=True)
        log = klasor / "gunluk.txt"
        env = dict(os.environ, PYTHONUNBUFFERED="1", PYTHONIOENCODING="utf-8",
                   FK_VERI=str(T.veri_klasoru()))
        env.pop("FK_LUCA_CEREZ", None)
        cerez = _LUCA_OTURUM["cerez"] if time.time() - _LUCA_OTURUM["zaman"] < LUCA_OTURUM_OMRU else []
        yol = None
        if cerez:
            import tempfile
            fd, yol = tempfile.mkstemp(prefix="fk_luca_", suffix=".json")
            with os.fdopen(fd, "w", encoding="utf-8") as fh0:
                json.dump(cerez, fh0)
            try:
                os.chmod(yol, 0o600)
            except OSError:
                pass
            env["FK_LUCA_CEREZ"] = yol
        bayrak = 0x08000000 if sys.platform == "win32" else 0      # CREATE_NO_WINDOW: siyah pencere açılmasın
        exe = sys.executable
        if exe.lower().endswith("pythonw.exe"):
            exe = exe[:-len("pythonw.exe")] + "python.exe"
        fh = open(log, "a", encoding="utf-8")
        proc = subprocess.Popen([exe, "-m", MODUL, "--sessiz", "--cikti", str(klasor), *args],
                                stdout=fh, stderr=subprocess.STDOUT, stdin=subprocess.DEVNULL,
                                env=env, cwd=os.getcwd(), creationflags=bayrak)
        if yol:                       # otomasyon okuyup siler; süreç hiç okumazsa da 2 dk sonra silinir
            t0 = threading.Timer(120, lambda y=yol: os.path.exists(y) and os.remove(y)); t0.daemon = True; t0.start()
        _is.clear()
        _is.update(tur=tur, proc=proc, klasor=klasor, log=log, fh=fh,
                   basladi=datetime.now().isoformat(timespec="seconds"))


AY = r"20\d\d-(0[1-9]|1[0-2])"


def _klasor(ad: str) -> Path:
    if not re.fullmatch(AY, ad):
        raise HTTPException(404)
    k = _toplu_kok() / ad
    if not k.is_dir():
        raise HTTPException(404, "Çalışma bulunamadı")
    return k


def _json(y: Path) -> dict:
    try:
        return json.loads(y.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _son_satirlar(y: Path, n=12) -> list[str]:
    try:
        satirlar = y.read_text(encoding="utf-8", errors="replace").splitlines()
    except OSError:
        return []
    return [s for s in satirlar if s.strip()][-n:]


# ------------------------------------------------------------------ firma listesi
class FirmaGuncelle(BaseModel):
    firmalar: list[dict]          # [{"id", "dahil", "not"}]


@router.get("/firmalar")
def firmalar():
    d = T.defter_oku()
    return {"firmalar": [{**f, "defter": d.get(f["id"], "")} for f in T.liste_oku()], "dosya": str(T.liste_yolu())}


@router.put("/firmalar")
def firmalar_kaydet(g: FirmaGuncelle):
    mevcut = {f["id"]: f for f in T.liste_oku()}
    for f in g.firmalar:
        if f.get("id") in mevcut:
            mevcut[f["id"]]["dahil"] = bool(f.get("dahil"))
            if "not" in f:
                mevcut[f["id"]]["not"] = str(f.get("not") or "")
    defter = {f["id"]: str(f.get("defter") or "") for f in g.firmalar if f.get("id") in mevcut and "defter" in f}
    if defter:
        T.defter_yaz(defter)
    try:
        _liste_kaydet(list(mevcut.values()))
    except PermissionError:
        raise HTTPException(409, "firmalar.xlsx Excel'de açık; kapatıp tekrar deneyin.")
    return {"tamam": True}


def _liste_kaydet(firmalar: list[dict]):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    yol = T.liste_yolu()
    wb = Workbook(); ws = wb.active; ws.title = "Firmalar"
    ws.append(T.BASLIKLAR)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3864")
    for i, f in enumerate(firmalar, 1):
        ws.append([i, f["ad"], f["id"], "E" if f["dahil"] else "H", f.get("not", "")])
    for col, w in zip("ABCDE", (6, 30, 14, 12, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    wb.save(yol)


class MuhaWinFirmalari(BaseModel):
    firmalar: list[dict]          # [{"id": SIRKET_ID, "ad": Luca kısa adı}]
    tam: bool = False             # True: Luca'nın tüm listesi (listede olmayanlar çıkar); False: yalnız ekle/güncelle


@router.post("/firmalar/muhawin")
def firmalar_muhawin(g: MuhaWinFirmalari):
    """MuhaWin'in Luca bağlantısından gelen firmalar (ana programdaki firmalar ya da 'Luca'dan listeyi çek')."""
    defter = {str(f.get("id")): str(f.get("defter") or "") for f in g.firmalar if "defter" in f and f.get("id")}
    if defter:
        T.defter_yaz(defter)
    gelen = []
    for f in g.firmalar:
        fid, ad = str(f.get("id") or "").strip(), str(f.get("ad") or "").strip()
        if fid and ad and fid not in {x["id"] for x in gelen}:
            gelen.append({"id": fid, "ad": ad})
    if g.tam:
        yeni = gelen
    else:
        mevcut = T.liste_oku()
        adlar = {f["id"]: f["ad"] for f in gelen}
        yeni = [{"id": f["id"], "ad": adlar.get(f["id"], f["ad"])} for f in mevcut]
        bilinen = {f["id"] for f in mevcut}
        yeni += [f for f in gelen if f["id"] not in bilinen]
        if [(f["id"], f["ad"]) for f in yeni] == [(f["id"], f["ad"]) for f in mevcut]:
            return {"tamam": True, "eklenen": 0, "toplam": len(mevcut)}
    eski = {f["id"] for f in T.liste_oku()}
    try:
        T.liste_yaz(yeni)
    except PermissionError:
        raise HTTPException(409, "firmalar.xlsx Excel'de açık; kapatıp tekrar deneyin.")
    return {"tamam": True, "eklenen": sum(f["id"] not in eski for f in yeni), "toplam": len(yeni)}


@router.post("/firmalar/cek")
def firmalar_cek():
    _baslat("liste", ["--liste"], _toplu_kok() / "_liste")
    return {"tamam": True}


# ------------------------------------------------------------------ çalıştırma
class Baslat(BaseModel):
    ay: str                         # YYYY-AA
    idler: list[str] | None = None  # boşsa Dahil=E olanlar
    klasor: str | None = None       # verilirse o çalışmaya eklenir (yeniden çalıştırma)
    adim: str | None = None         # verilirse yalnız bu adım tekrar denenir (tek firma, mevcut çalışma)
    yalniz_fatura: bool = False     # Tek firma › "Luca'dan al": yalnız e-belgeler, muavin korunur, toplu özetine girmez


@router.post("/baslat")
def baslat(b: Baslat):
    if not re.fullmatch(r"20\d\d-(0[1-9]|1[0-2])", b.ay):
        raise HTTPException(400, "Ay YYYY-AA biçiminde olmalı")
    firmalar = T.liste_oku()
    if not firmalar:
        raise HTTPException(400, "Firma listesi yok. Önce 'Luca'dan firma listesini çek' düğmesine basın.")
    args = ["--ay", b.ay]
    if b.idler:
        args += ["--idler", ",".join(b.idler)]
    elif not any(f["dahil"] for f in firmalar):
        raise HTTPException(400, "Dahil işaretli firma yok.")
    if b.yalniz_fatura:
        if not b.idler or len(b.idler) != 1:
            raise HTTPException(400, "Luca'dan fatura almak için tek firma seçin")
        if b.idler[0] not in {f["id"] for f in firmalar}:
            raise HTTPException(400, "Firma toplu kontrol listesinde yok. MuhaWin'de firmayı Luca'dan ekleyin.")
        if b.adim:                  # yalnız indirilemeyen tür (ya da Denetim › "Muavini Luca'dan Al": muavin) alınır; diğerleri korunur
            if b.adim not in T.ADIM_KODLARI:
                raise HTTPException(400, "Bilinmeyen belge türü")
            args += ["--adim", b.adim]
        k = _toplu_kok() / "_tek" / b.ay
        _baslat("toplu", args + ["--yalniz-fatura"], k)
        return {"klasor": "_tek/" + b.ay}
    if b.adim:
        if b.adim not in T.ADIM_KODLARI or not b.klasor or not b.idler or len(b.idler) != 1:
            raise HTTPException(400, "Adım tekrarı için tek firma ve mevcut çalışma gerekir")
        args += ["--adim", b.adim]
    if b.klasor:
        if b.adim == "muavin" and re.fullmatch(AY, b.klasor):
            (_toplu_kok() / b.klasor).mkdir(parents=True, exist_ok=True)   # yalnız muavin: önceki çalışma gerekmez
        k = _klasor(b.klasor)
        if k.name != b.ay:
            raise HTTPException(400, "Çalışma ile ay uyuşmuyor")
    # her ayın tek bir özeti vardır; yeniden çalıştırılan firmaların sonuçları eskisinin yerine geçer
    k = _toplu_kok() / b.ay
    _baslat("toplu", args, k)
    return {"klasor": k.name}


@router.post("/durdur")
def durdur():
    p = _is.get("proc")
    if not _calisiyor():
        return {"tamam": True}
    if sys.platform == "win32":
        subprocess.run(["taskkill", "/PID", str(p.pid), "/T", "/F"], capture_output=True,
                       creationflags=0x08000000)
    else:
        p.terminate()
    try:
        p.wait(10)
    except subprocess.TimeoutExpired:
        p.kill()
    return {"tamam": True}


@router.get("/durum")
def durum():
    if not _is:
        return {"calisiyor": False}
    k: Path = _is["klasor"]
    d = _json(k / "durum.json")
    calisiyor = _calisiyor()
    kod = _is["proc"].poll()
    if not calisiyor and _is.get("fh"):
        _is["fh"].close(); _is["fh"] = None
    sessiz_sn = None
    if d.get("guncelleme"):
        sessiz_sn = (datetime.now() - datetime.fromisoformat(d["guncelleme"])).total_seconds()
    cr = k / "captcha.png"
    captcha = int(cr.stat().st_mtime * 1000) if calisiyor and cr.exists() else 0
    return {"calisiyor": calisiyor, "tur": _is["tur"], "klasor": k.name, "basladi": _is["basladi"],
            "cikis_kodu": kod, "ilerleme": d, "sessiz_sn": sessiz_sn, "captcha": captcha,
            "gunluk": _son_satirlar(_is["log"]),
            "sonuclar": _adimli(k, _json(k / "ozet.json").get("sonuclar", [])) if _is["tur"] == "toplu" else []}


# ------------------------------------------------------------------ geçmiş
# Luca doğrulama kodu (captcha) programda: otomasyonun kaydettiği resim gösterilir, kullanıcının yazdığı kod otomasyona verilir
@router.get("/captcha")
def captcha_resmi():
    from fastapi.responses import Response
    y = (_is.get("klasor") or Path("/yok")) / "captcha.png"
    if not _calisiyor() or not y.exists():
        raise HTTPException(404, "Doğrulama kodu beklenmiyor")
    return Response(y.read_bytes(), media_type="image/png", headers={"Cache-Control": "no-store"})


class CaptchaYanit(BaseModel):
    kod: str


@router.post("/captcha")
def captcha_yanit(b: CaptchaYanit):
    kod = re.sub(r"\s+", "", b.kod or "")[:20]
    if not kod:
        raise HTTPException(400, "Kodu yazın")
    k = _is.get("klasor")
    if not k or not _calisiyor() or not (k / "captcha.png").exists():
        raise HTTPException(409, "Doğrulama kodu beklenmiyor")
    (k / "captcha_yanit.txt").write_text(kod, encoding="utf-8")
    (k / "captcha.png").unlink(missing_ok=True)
    return {"tamam": True}


@router.get("/gecmis")
def gecmis():
    kok = _toplu_kok()
    cikti = []
    for k in sorted(kok.iterdir(), reverse=True) if kok.exists() else []:
        if not k.is_dir() or not re.fullmatch(AY, k.name):
            continue
        oz = _json(k / "ozet.json")
        s = oz.get("sonuclar", [])
        say = lambda d: sum(r.get("durum") == d for r in s)
        y = k / "ozet.json"
        cikti.append({"klasor": k.name, "ay": k.name,
                      "zaman": datetime.fromtimestamp(y.stat().st_mtime).isoformat(timespec="minutes")
                      if y.exists() else "",
                      "firma": len(s), "temiz": say("Temiz"), "sorunlu": say("Sorun var"),
                      "hata": say("Hata") + say("Eksik indirme"), "belgesiz": say("Belge/muavin yok"),
                      "ozet_var": (k / "toplu_ozet.xlsx").exists()})
    return cikti


def _adimli(k: Path, sonuclar: list[dict]) -> list[dict]:
    """Satırda adım özeti yoksa firmanın dönem klasöründeki adimlar.json'dan üret."""
    bas = datetime.strptime(k.name + "-01", "%Y-%m-%d").date()
    for r in sonuclar:
        if "adimlar" not in r:
            r["adimlar"] = T.adim_ozeti(T.adimlari_oku(T.firma_klasoru(bas, r)))
    return sonuclar


def _guncelle(sonuclar: list[dict]) -> bool:
    from ..rapor_kayit import KURAL_SURUMU, rapor_oku
    degisti = False
    for r in sonuclar:
        rid = r.get("rapor_id")
        if not rid or r.get("kural") == KURAL_SURUMU:
            continue
        v = rapor_oku(rid)
        if not v or v.get("kural") != KURAL_SURUMU:
            continue
        r["sorunlar"] = {d["ad"]: d["adet"] for d in v["durumlar"]}
        r["kural"] = KURAL_SURUMU
        if r.get("durum") in ("Sorun var", "Temiz"):
            r["durum"] = "Sorun var" if sum(r["sorunlar"].get(x, 0) for x in T.SORUN_SUTUNLARI) else "Temiz"
        degisti = True
    return degisti


@router.get("/calisma/{ad}")
def calisma(ad: str):
    k = _klasor(ad)
    oz = _json(k / "ozet.json")
    if _guncelle(oz.get("sonuclar", [])):              # kurallar değiştiyse sayılar raporlardan yenilenir
        try:
            (k / "ozet.json").write_text(json.dumps(oz, ensure_ascii=False, indent=1), encoding="utf-8")
        except OSError:
            pass
    return {"klasor": ad, "bas": oz.get("bas"), "bit": oz.get("bit"), "sonuclar": _adimli(k, oz.get("sonuclar", [])),
            "ozet_var": (k / "toplu_ozet.xlsx").exists()}


class LucaOturum(BaseModel):
    cerezler: list[dict] = []


_CEREZ_ALAN = {"name", "value", "domain", "path", "url", "expires", "httpOnly", "secure", "sameSite"}


@router.post("/luca-oturum")
def luca_oturum(b: LucaOturum):
    """MuhaWin, Luca'ya girildiğinde oturum çerezlerini bildirir (boş liste: oturum kapandı). Yalnız *.luca.com.tr."""
    temiz = []
    for c in b.cerezler[:300]:
        c = {k: v for k, v in c.items() if k in _CEREZ_ALAN}
        yer = str(c.get("domain") or c.get("url") or "").lower()
        host = re.sub(r"^https?://", "", yer).split("/")[0].lstrip(".")
        if c.get("name") and isinstance(c.get("value"), str) and (host == "luca.com.tr" or host.endswith(".luca.com.tr")):
            temiz.append(c)
    _LUCA_OTURUM.update(cerez=temiz, zaman=time.time())
    return {"tamam": True, "adet": len(temiz)}


class SonucSil(BaseModel):
    anahtarlar: list[str]           # firma id'leri (id'siz eski satırlar için firma adı)
    belgeler: bool = False          # firmanın o aya ait indirilen belgeleri ve kontrol raporu da silinsin


@router.post("/calisma/{ad}/sil")
def calisma_sonuc_sil(ad: str, b: SonucSil):
    """Seçilen firmaların bu aydaki toplu kontrol sonuçlarını özetten çıkarır; istenirse dönem klasörünü de siler."""
    import shutil
    if _calisiyor():
        raise HTTPException(409, "Çalışan bir işlem var. Bitmesini bekleyin ya da durdurun.")
    k = _klasor(ad)
    oz = _json(k / "ozet.json")
    sec = set(b.anahtarlar or [])
    if not sec:
        raise HTTPException(400, "Firma seçilmedi")
    kalan, silinen = [], []
    for r in oz.get("sonuclar", []):
        (silinen if (r.get("id") or r.get("ad")) in sec else kalan).append(r)
    if not silinen:
        raise HTTPException(404, "Seçilen firmaların sonucu bu ayda yok")
    bas = date.fromisoformat(oz.get("bas") or ad + "-01")
    bit = date.fromisoformat(oz["bit"]) if oz.get("bit") else bas
    klasor_silinen = 0
    if b.belgeler:
        kok = depo.firmalar_kok().resolve()
        for r in silinen:
            fk = T.firma_klasoru(bas, r)
            try:
                if fk.is_dir() and kok in fk.resolve().parents:
                    shutil.rmtree(fk); klasor_silinen += 1
            except OSError as e:
                raise HTTPException(409, f"{r.get('ad')} klasörü silinemedi (dosya açık olabilir): {e}")
    try:
        T.ozet_yaz(kalan, bas, bit, k, birlestir=False)
    except PermissionError:
        raise HTTPException(409, "toplu_ozet.xlsx Excel'de açık; kapatıp tekrar deneyin.")
    return {"silinen": len(silinen), "klasor_silinen": klasor_silinen, "kalan": len(kalan)}


@router.get("/calisma/{ad}/kdv2")
def calisma_kdv2(ad: str):
    k = _klasor(ad)
    oz = _json(k / "ozet.json")
    if _guncelle(oz.get("sonuclar", [])):              # kurallar değiştiyse sayılar raporlardan yenilenir
        try:
            (k / "ozet.json").write_text(json.dumps(oz, ensure_ascii=False, indent=1), encoding="utf-8")
        except OSError:
            pass
    return {"klasor": ad, "bas": oz.get("bas"), "bit": oz.get("bit"),
            "firmalar": T.kdv2_raporu(oz.get("sonuclar", []), oz.get("bas"), oz.get("bit"))}


@router.get("/calisma/{ad}/kdv2.xlsx")
def calisma_kdv2_excel(ad: str):
    from openpyxl import Workbook
    k = _klasor(ad)
    oz = _json(k / "ozet.json")
    wb = Workbook()
    wb.remove(wb.active)
    ay = f"{ad[5:7]}.{ad[:4]}"
    T.kdv2_sayfasi(wb, oz.get("sonuclar", []), f"KDV2 Raporu – {ay} tevkifatlı alış faturaları",
                   oz.get("bas"), oz.get("bit"))
    y = k / "kdv2_raporu.xlsx"
    try:
        wb.save(y)
    except PermissionError:
        raise HTTPException(409, "kdv2_raporu.xlsx Excel'de açık; kapatıp tekrar deneyin.")
    return FileResponse(y, filename=f"KDV2_raporu_{ad}.xlsx")


@router.get("/calisma/{ad}/ozet.xlsx")
def calisma_ozet(ad: str):
    y = _klasor(ad) / "toplu_ozet.xlsx"
    if not y.exists():
        raise HTTPException(404, "Özet henüz yok")
    return FileResponse(y, filename=f"toplu_ozet_{ad}.xlsx")


@router.post("/calisma/{ad}/klasoru-ac")
def klasoru_ac(ad: str):
    k = _klasor(ad)
    if sys.platform == "win32":
        os.startfile(k)     # noqa
    return {"klasor": str(k)}


@router.post("/firma-klasoru")
def firma_klasoru_ac(b: dict):
    """Firmanın o dönemdeki belgelerinin bulunduğu klasörü Windows Gezgini'nde açar."""
    k = depo.rapor_klasoru(str(b.get("id") or ""))
    if not k or not k.exists():
        raise HTTPException(404, "Klasör bulunamadı")
    if sys.platform == "win32":
        os.startfile(k)     # noqa
    return {"klasor": str(k)}
