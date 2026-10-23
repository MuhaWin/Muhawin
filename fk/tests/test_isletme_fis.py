"""İşletme / serbest meslek defteri: faturalardan Luca "Excel'den fiş aktarma" dosyası.
Örnek: kullanıcının gönderdiği iki e-Arşiv faturası (MTMX, Yalçın Erten) ve onlara göre doldurulmuş ÖRNEK FİŞ EXCEL."""
import io
import zipfile
from datetime import date
from pathlib import Path

import pytest

FIRMA_VKN, FIRMA_TCKN = "2790296356", "19006752134"     # Olympos Quad Safari (Murat Demiralp) — şahıs firması


def _taraf(vkn, sema, ad, vd, sokak, ilce, il, kisi=None):
    kisi_xml = f"<cac:Person><cbc:FirstName>{kisi[0]}</cbc:FirstName><cbc:FamilyName>{kisi[1]}</cbc:FamilyName></cac:Person>" if kisi else ""
    ad_xml = f"<cac:PartyName><cbc:Name>{ad}</cbc:Name></cac:PartyName>" if ad else ""
    return f"""<cac:Party><cac:PartyIdentification><cbc:ID schemeID="{sema}">{vkn}</cbc:ID></cac:PartyIdentification>{ad_xml}
<cac:PostalAddress><cbc:StreetName>{sokak}</cbc:StreetName><cbc:CitySubdivisionName>{ilce}</cbc:CitySubdivisionName><cbc:CityName>{il}</cbc:CityName></cac:PostalAddress>
<cac:PartyTaxScheme><cac:TaxScheme><cbc:Name>{vd}</cbc:Name></cac:TaxScheme></cac:PartyTaxScheme>{kisi_xml}</cac:Party>"""


def _fatura(no, tarih, satici, alici, satirlar, matrah, kdv, tip="SATIS"):
    s = "".join(f"""<cac:InvoiceLine><cbc:ID>{i}</cbc:ID><cbc:InvoicedQuantity unitCode="C62">{m}</cbc:InvoicedQuantity>
<cbc:LineExtensionAmount currencyID="TRY">{t}</cbc:LineExtensionAmount><cac:TaxTotal><cbc:TaxAmount currencyID="TRY">{k}</cbc:TaxAmount>
<cac:TaxSubtotal><cbc:TaxableAmount currencyID="TRY">{t}</cbc:TaxableAmount><cbc:TaxAmount currencyID="TRY">{k}</cbc:TaxAmount><cbc:Percent>20</cbc:Percent>
<cac:TaxCategory><cac:TaxScheme><cbc:TaxTypeCode>0015</cbc:TaxTypeCode></cac:TaxScheme></cac:TaxCategory></cac:TaxSubtotal></cac:TaxTotal>
<cac:Item><cbc:Name>{ad}</cbc:Name></cac:Item></cac:InvoiceLine>""" for i, (ad, m, t, k) in enumerate(satirlar, 1))
    top = f"{float(matrah) + float(kdv):.2f}"
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<Invoice xmlns="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2"
 xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2"
 xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2">
<cbc:ProfileID>EARSIVFATURA</cbc:ProfileID><cbc:ID>{no}</cbc:ID><cbc:UUID>{no}-uuid</cbc:UUID>
<cbc:IssueDate>{tarih}</cbc:IssueDate><cbc:InvoiceTypeCode>{tip}</cbc:InvoiceTypeCode><cbc:DocumentCurrencyCode>TRY</cbc:DocumentCurrencyCode>
<cac:AccountingSupplierParty>{satici}</cac:AccountingSupplierParty><cac:AccountingCustomerParty>{alici}</cac:AccountingCustomerParty>
<cac:TaxTotal><cbc:TaxAmount currencyID="TRY">{kdv}</cbc:TaxAmount><cac:TaxSubtotal>
<cbc:TaxableAmount currencyID="TRY">{matrah}</cbc:TaxableAmount><cbc:TaxAmount currencyID="TRY">{kdv}</cbc:TaxAmount><cbc:Percent>20</cbc:Percent>
<cac:TaxCategory><cac:TaxScheme><cbc:TaxTypeCode>0015</cbc:TaxTypeCode></cac:TaxScheme></cac:TaxCategory></cac:TaxSubtotal></cac:TaxTotal>
{s}<cac:LegalMonetaryTotal><cbc:LineExtensionAmount currencyID="TRY">{matrah}</cbc:LineExtensionAmount><cbc:TaxExclusiveAmount currencyID="TRY">{matrah}</cbc:TaxExclusiveAmount>
<cbc:TaxInclusiveAmount currencyID="TRY">{top}</cbc:TaxInclusiveAmount><cbc:PayableAmount currencyID="TRY">{top}</cbc:PayableAmount></cac:LegalMonetaryTotal>
</Invoice>""".encode()


def ornek_zip(yol: Path) -> Path:
    mtmx = _fatura("MIT2026000036281", "2026-09-01",
                   _taraf("4780120398", "VKN", "MTMX PAZARLAMA VE TİCARET A.Ş", "KADIKÖY VERGİ DAİRESİ", "Eğitim Mah.Poyraz Sk Sadıkoğlu 2 İş Merkezi No:6", "KADIKÖY", "İstanbul"),
                   _taraf(FIRMA_VKN, "VKN", "Olympos Quad Safari", "KUMLUCA VERGİ DAİRESİ", "Yazır mah. Olympos sok. No:5", "KUMLUCA", "Antalya"),
                   [("Just1 J-STR Kapalı Motosiklet Kaskı", "1", "2683.33", "536.67"), ("KARGO BEDELİ", "1", "140.83", "28.16")], "2824.17", "564.83")
    erten = _fatura("N012026000001409", "2026-09-04",
                    _taraf("11729030070", "TCKN", "", "SARAYLAR", "hacıkaplanlar mah. 750 sk", "Pamukkale", "Denizli", ("YALÇIN", "ERTEN")),
                    _taraf(FIRMA_TCKN, "TCKN", "", "", "YAZIR MAH OLYMPOS SOK NO 5", "Kumluca", "Antalya", ("MURAT", "DEMİRALP")),
                    [("SYM ATV ALTERNATÖR", "1", "5500.00", "1100.00")], "5500.00", "1100.00")
    with zipfile.ZipFile(yol, "w") as z:
        z.writestr("a.xml", mtmx); z.writestr("b.xml", erten)
    return yol


@pytest.fixture
def veri(tmp_path, monkeypatch):
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    return tmp_path


def test_ornek_excel_ile_ayni(veri):
    from fatura_kontrol import isletme_fis as I
    z = ornek_zip(veri / "f.zip")
    I.ayar_yaz("OLYMPOS", {"faaliyet_kodu": "MOTOSİKLET, KARAVAN VE KAMP GEREÇLERİNİN OPERATÖRSÜZ OLARAK KİRALANMASI VEYA LEASİNGİ (FİNANSAL LEASİNG HARİÇ)773903"})
    s = I.satirlar_hazirla([z], "OLYMPOS", {FIRMA_VKN})
    assert [d["belge_no"] for d in s["disarida"]] == ["N012026000001409"]        # TCKN'li fatura: firma kimliği eksik
    I.ayar_yaz("OLYMPOS", {"kimlikler": [FIRMA_TCKN]})
    s = I.satirlar_hazirla([z], "OLYMPOS", {FIRMA_VKN})
    assert not s["disarida"] and len(s["satirlar"]) == 2
    a, b = s["satirlar"]
    assert (a["islem"], a["kategori"], a["belge_turu"], a["evrak_tarihi"], a["kayit_tarihi"], a["evrak_no"], a["vkn"]) == \
        ("Gider", "Defter Fişleri", "Diğer Alışlar", "01/09/2026", "01/09/2026", "MIT2026000036281", "4780120398")
    assert a["vergi_dairesi"] == "034272  KADIKÖY VERGİ DAİRESİ"
    assert (a["unvan1"], a["unvan2"]) == ("MTMX PAZARLAMA VE Tİ", "CARET ANONİM ŞİRKETİ")      # örnek Excel'deki bölme
    assert (a["belge_turu_db"], a["alis_satis_turu"], a["aciklama"]) == ("e-Arşiv Fatura", "Normal Alım", "MTMX PAZARLAMA VE TİCARET A.Ş")
    assert (a["miktar"], a["birim_fiyat"], a["tutar"], a["tevkifat"], a["kdv_orani"], a["kdv_tutari"], a["toplam"]) == \
        (2.0, 1412.09, 2824.17, "Yok", 20, 564.83, 3389.0)
    assert a["faaliyet_kodu"].endswith("773903")
    assert (b["vkn"], b["unvan1"], b["unvan2"], b["vergi_dairesi"]) == ("11729030070", "ERTEN", "YALÇIN", "020201  SARAYLAR VERGİ DAİRESİ")
    assert (b["miktar"], b["birim_fiyat"], b["tutar"], b["kdv_tutari"], b["toplam"]) == (1.0, 5500.0, 5500.0, 1100.0, 6600.0)
    assert b["adres"].startswith("hacıkaplanlar mah. 750 sk")
    # Excel: örnek dosyayla aynı başlıklar ve hücre türleri
    import openpyxl
    y = I.excel_yaz(s["satirlar"], veri / "fis.xlsx")
    ws = openpyxl.load_workbook(y).active
    assert [c.value for c in ws[1]] == I.BASLIKLAR and len(I.BASLIKLAR) == 33
    r = {I.BASLIKLAR[i]: c.value for i, c in enumerate(ws[2])}
    assert r["EVRAK TARİHİ"] == "01/09/2026" and r["TUTAR"] == 2824.17 and r["KDV ORANI"] == 20 and r["TOPLAM TUTAR"] == 3389
    assert r["KREDİLİ TUTAR"] == 0 and r["STOPAJ TUTARI"] == 0 and r["SERİ NO"] in ("", None)


def test_secimler_ogrenilir_ve_kisaltmalar(veri):
    from fatura_kontrol import isletme_fis as I
    z = ornek_zip(veri / "f.zip")
    s = I.satirlar_hazirla([z], "OLYMPOS", {FIRMA_VKN, FIRMA_TCKN})
    for r in s["satirlar"]:
        r["kayit_alt_turu"] = "4-185" if r["vkn"] == "4780120398" else "4-114"
    I.secimleri_ogren("OLYMPOS", s["satirlar"])
    s2 = I.satirlar_hazirla([z], "OLYMPOS", {FIRMA_VKN, FIRMA_TCKN})
    assert [r["kayit_alt_turu"] for r in s2["satirlar"]] == ["4-185", "4-114"] and all(r["_ogrenildi"] for r in s2["satirlar"])
    assert I.kisaltmalari_ac("Expo İnş. San. ve Tic. Ltd. Şti.") == "EXPO İNŞAAT SANAYİ VE TİCARET LİMİTED ŞİRKETİ"
    assert I.kisaltmalari_ac("ABC GIDA A.Ş.") == "ABC GIDA ANONİM ŞİRKETİ"
    assert I.kisaltmalari_ac("SANAT EVİ") == "SANAT EVİ"                 # sözcük içindeki SAN açılmaz


def test_luca_listeleri(veri):
    from fatura_kontrol import isletme_fis as I
    r = I.listeleri_kaydet([
        {"ad": "kayitAltTuru", "etiket": "Kayıt Alt Türü", "secenekler": ["Seçiniz", "4-114", "4-185"]},
        {"ad": "vergiDairesi", "etiket": "Vergi Dairesi", "secenekler": ["034272  KADIKÖY VERGİ DAİRESİ", "007201  KUMLUCA VERGİ DAİRESİ"]},
        {"ad": "belgeTuru", "etiket": "Belge Türü", "secenekler": ["Diğer Alışlar", "Diğer Satışlar", "Satış"]}], "2026-10-22")
    assert r["bulunan"] == ["belge_turu", "kayit_alt_turu", "vergi_dairesi"]
    L = I.listeler_oku()
    assert L["kayit_alt_turu"] == ["4-114", "4-185"] and L["_kaynak"]["luca"]
    assert I.vergi_dairesi("Kumluca", "Antalya", L) == "007201  KUMLUCA VERGİ DAİRESİ"
    assert I.vergi_dairesi("Kumluca", "Ankara", L) == "KUMLUCA VERGİ DAİRESİ"         # başka ilde: kod uydurulmaz


def test_isletme_aktarimi_simdilik_kapali(veri):
    """Kullanıcı kararı (23.10.2026): işletme / serbest meslek fatura aktarımı "Çok yakında"; sayfa ve Excel kapalı."""
    from fastapi.testclient import TestClient
    from fatura_kontrol.web import app as A
    c = TestClient(A.app)
    assert not A.ISLETME_AKTARIM_ACIK and "Çok yakında" in c.get("/isletme-fis").text
    assert c.post("/api/isletme-fis/excel", json={"id": "X/2026/09", "satirlar": []}).status_code == 403


def test_api_ve_sayfa(veri, monkeypatch):
    import shutil
    from fastapi.testclient import TestClient
    from fatura_kontrol import depo
    from fatura_kontrol.web import app as A
    monkeypatch.setattr(A, "ISLETME_AKTARIM_ACIK", True)          # özellik açıldığında çalışmaya hazır kalsın
    c = TestClient(A.app)
    k = depo.donem_klasoru("OLYMPOS", date(2026, 9, 1)); k.mkdir(parents=True)
    shutil.copyfile(ornek_zip(veri / "f.zip"), k / "faturalar.zip")
    depo.firma_bilgisi_yaz("OLYMPOS", FIRMA_VKN, "555")
    assert "Luca'ya Excel'den fiş aktarma" in c.get("/isletme-fis").text
    assert c.get("/api/isletme-fis", params={"firma_id": "555", "bas": "2026-10-01"}).status_code == 400   # o ay fatura yok
    v = c.get("/api/isletme-fis", params={"id": "OLYMPOS/2026/09"}).json()
    assert len(v["satirlar"]) == 1 and len(v["disarida"]) == 1 and v["listeler"]["belge_turu"]
    assert c.put("/api/isletme-fis/ayar", json={"firma": "OLYMPOS", "kimlikler": [FIRMA_TCKN], "kategori": "Serbest Meslek Fişleri"}).status_code == 200
    v = c.get("/api/isletme-fis", params={"firma_id": "555", "bas": "2026-09-15"}).json()
    assert len(v["satirlar"]) == 2 and all(r["kategori"] == "Serbest Meslek Fişleri" for r in v["satirlar"])
    for r in v["satirlar"]:
        r["kayit_alt_turu"] = "4-185"
    assert c.post("/api/isletme-fis/excel", json={"id": "OLYMPOS/2026/09", "satirlar": [{**v["satirlar"][0], "tutar": "abc"}]}).status_code == 400
    j = c.post("/api/isletme-fis/excel", json={"id": "OLYMPOS/2026/09", "satirlar": v["satirlar"]}).json()
    assert j["adet"] == 2 and (k / "luca_fis_aktarim.xlsx").exists()
    d = c.get(j["indir"]); assert d.status_code == 200 and d.content[:2] == b"PK"
    v2 = c.get("/api/isletme-fis", params={"id": "OLYMPOS/2026/09"}).json()
    assert all(r["kayit_alt_turu"] == "4-185" and r["_ogrenildi"] for r in v2["satirlar"])
    assert c.post("/api/isletme-fis/listeler", json={"listeler": []}).status_code == 400
    r = c.post("/api/isletme-fis/listeler", json={"listeler": [{"ad": "altTur", "etiket": "Kayıt Alt Türü", "secenekler": ["4-114", "4-185"]}]}).json()
    assert r["bulunan"] == ["kayit_alt_turu"] and c.get("/api/isletme-fis/listeler").json()["kayit_alt_turu"] == ["4-114", "4-185"]


def test_defter_turu_ile_dugme(veri):
    """'Luca fiş Excel'i (işletme)' düğmesi yalnız işletme / serbest meslek defterinde: rapor JSON'unda defter türü."""
    from fatura_kontrol import depo
    from fatura_kontrol.luca import toplu as T
    from fatura_kontrol.web.app import defter_turu
    k = depo.donem_klasoru("OLYMPOS", date(2026, 9, 1)); k.mkdir(parents=True)
    assert defter_turu(k) == ""
    depo.firma_bilgisi_yaz("OLYMPOS", FIRMA_VKN, "555")
    T.defter_yaz({"555": "serbest"})
    assert defter_turu(k) == "serbest"
    T.defter_yaz({"555": "bilanco"})
    assert defter_turu(k) == "bilanco"
    k2 = depo.donem_klasoru("BAŞKA", date(2026, 9, 1)); k2.mkdir(parents=True)
    from fatura_kontrol import muavin_programlari as MP
    MP.klasore_yaz(k2, "luca_isletme")
    assert defter_turu(k2) == "isletme"


def test_muavin_yukle_isletme_listesi(tmp_path, monkeypatch):
    """MuhaWin'den (Luca oturumuyla) gelen işletme / serbest meslek Gelir ve Gider Listesi dönem klasörüne yazılır."""
    import base64, json as _j
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fastapi.testclient import TestClient
    from fatura_kontrol.web import app as A
    from fatura_kontrol import depo
    from fatura_kontrol.kontrol import muavin_dosyalari
    from datetime import date as _d
    L = {"firma": "AV. AHMET", "bas": "2026-01-01", "bit": "2026-12-31", "gelir": [
        {"fid": "55", "tarih": "05/09/2026", "evrak_no": "SMM2026000000705", "kategori": "Serbest Meslek Fişleri", "belge_turu": "e-SMM",
         "aciklama": "Danışmanlık", "cari": "XYZ LTD", "tutar": "1.000,00", "kdv": "200,00", "toplam": "1.200,00"}], "gider": []}
    r = TestClient(A.app).post("/api/muavin/yukle", json={"ad": "AV. AHMET", "luca_id": "77", "yil": 2026, "ay": 9, "program": "luca_isletme",
                                                          "dosya_adi": "muavin_isletme.json", "b64": base64.b64encode(_j.dumps(L).encode()).decode()})
    assert r.status_code == 200, r.text
    k = depo.donem_klasoru("AV. AHMET", _d(2026, 9, 1))
    assert [y.name for y in muavin_dosyalari(k)] == ["muavin_isletme.json"] and r.json()["satir"] >= 1
