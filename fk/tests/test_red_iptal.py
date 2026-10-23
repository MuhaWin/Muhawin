"""Red / iptal fatura listeleri (Mikro e-fatura / e-arşiv Excel'i): kontrol, rapor, Excel ve çıktı raporu."""
import io
import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).parent))
from test_luca_otomasyon import FIRMA_VKN, muavin_yaz, ubl   # noqa: E402

BASLIK = ["Faturaya bağlı", "Fatura Tipi", "Belge tarihi", "EFatura ID", "Fatura türü", "Vergi kimlik no", "Cari kod", "Cari ünvanı",
          "Ara toplam", "Vergi matrahı", "Hesaplanan KDV", "İskonto", "Yekün", "Döviz Cinsi", "Satır Sayısı", "Fatura durum kodu",
          "Fatura durum açıklaması", "GİB'den gelen durum kodu", "Zarf durum açıklaması", "GİB durum açıklaması"]


def red_xlsx(yol: Path, satirlar):
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Eylül e fatura satış"
    ws.append(BASLIK)
    for no, seri_tarih, vkn, unvan, matrah, kdv, durum in satirlar:
        ws.append(["Yanlış", "Ticari", seri_tarih, no, "SATIS", vkn, "120.01", unvan, matrah, matrah, kdv, 0, matrah + kdv, "TRY", 1,
                   "2002", durum, "1300", "Zarf gönderildi", "BAŞARIYLA TAMAMLANDI"])
    wb.save(yol)
    return yol


@pytest.fixture
def c(tmp_path, monkeypatch):
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fastapi.testclient import TestClient
    from fatura_kontrol.web import app as A
    return TestClient(A.app)


def test_red_iptal_okuma(tmp_path):
    from datetime import date
    from fatura_kontrol import red_iptal as R
    y = red_xlsx(tmp_path / "Eylül e fatura satış.xlsx", [("aas2026000004610", 46267, "1180376519", "AYDOĞAN ENTEGRE A.Ş.", 2000, 400, "Fatura red edildi"),
                                                          ("AAS2026000004611", 46268, "1", "X LTD", 10, 2, "Fatura iptal edildi")])
    L, u = R.oku([y])
    assert not u and set(L) == {"AAS2026000004610", "AAS2026000004611"}
    a = L["AAS2026000004610"]
    assert a["tarih"] == date(2026, 9, 2) and a["vkn"] == "1180376519" and float(a["matrah"]) == 2000 and float(a["kdv"]) == 400
    assert a["durum"] == "Fatura red edildi" and a["red"] and a["yon"] == "SATIS"
    assert R.durum_metni(a) == "Red" and R.durum_metni(L["AAS2026000004611"]) == "İptal"
    bos = tmp_path / "bos.xlsx"
    import openpyxl; wb = openpyxl.Workbook(); wb.active.append(["a", "b"]); wb.save(bos)
    assert R.oku([bos])[1]


def test_red_iptal_kontrol_rapor_excel_cikti(c, tmp_path):
    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w") as z:
        z.writestr("a.xml", ubl("AAA2026000000001", "2026-09-05", "1", FIRMA_VKN, "1000.00", "200.00", 20))
        z.writestr("b.xml", ubl("EX02026000000001", "2026-09-10", FIRMA_VKN, "9", "2000.00", "400.00", 20))
        z.writestr("c.xml", ubl("ZZZ2026000000005", "2026-09-11", FIRMA_VKN, "7", "100.00", "20.00", 20))
    m = tmp_path / "muavin.xlsx"; muavin_yaz(m)
    r1 = red_xlsx(tmp_path / "Eylül e fatura satış.xls.xlsx", [
        ("EX02026000000001", 46275, "9", "ALICI 9", 2000, 400, "Fatura red edildi"),     # dosyada var, muavinde kayıtlı → sorun
        ("ZZZ2026000000005", 46276, "7", "ALICI 7", 100, 20, "Fatura red edildi")])     # dosyada var, kayıt yok → uyumlu
    r2 = red_xlsx(tmp_path / "Eylül alış ft.xlsx", [
        ("AAA2026000000002", 46271, "2", "SATICI 2", 750, 150, "Fatura red edildi"),     # dosyada yok, muavinde kayıtlı → sorun
        ("QQQ2026000000009", 46272, "8", "SATICI 8", 50, 10, "Fatura iptal edildi")])    # hiçbir yerde yok → uyumlu
    r = c.post("/api/kontrol", data={"bas": "2026-09-01", "bit": "2026-09-30", "program": "luca"},
               files=[("faturalar", ("p.zip", zb.getvalue(), "application/zip")), ("muavinler", ("muavin.xlsx", m.read_bytes(), "application/octet-stream")),
                      ("red_iptal", (r1.name, r1.read_bytes(), "application/octet-stream")), ("red_iptal", (r2.name, r2.read_bytes(), "application/octet-stream"))])
    assert r.status_code == 200, r.text
    v = r.json()
    d = {s["belge_no"]: s["durum"] for s in v["sonuclar"]}
    assert d["EX02026000000001"] == "Red / iptal edilmiş fatura muavinde kayıtlı" and d["ZZZ2026000000005"] == "Red / iptal – muavinde kaydı yok, uyumlu"
    assert d["AAA2026000000001"] == "Uyumlu"
    red = {x["belge_no"]: x for x in v["red_iptal"]}
    assert set(red) == {"EX02026000000001", "ZZZ2026000000005", "AAA2026000000002", "QQQ2026000000009"}
    assert red["AAA2026000000002"]["muavinde"] and not red["AAA2026000000002"]["faturada"] and red["AAA2026000000002"]["yon"] == "Alış"
    assert not red["QQQ2026000000009"]["muavinde"] and red["QQQ2026000000009"]["ne"] == "İptal"
    assert "AAA2026000000002" not in {x["belge_no"] for x in v["faturasi_olmayanlar"]}     # red listesinde gösterilir
    say = {x["ad"]: x["adet"] for x in v["durumlar"]}
    assert say["Red / iptal edilmiş fatura muavinde kayıtlı"] == 2 and say["Red / iptal – muavinde kaydı yok, uyumlu"] == 2
    from fatura_kontrol import depo
    k = depo.rapor_klasoru(v["id"])
    assert len(list(k.glob("red_iptal_*"))) == 2
    import openpyxl
    wb = openpyxl.load_workbook(k / depo.RAPOR_EXCEL)
    assert "Red - İptal" in wb.sheetnames and wb["Red - İptal"].max_row == 5
    cikti = c.get("/api/rapor/cikti", params={"id": v["id"]}).text
    assert "Red / iptal edilmiş faturalar" in cikti and "2 tanesi muavinde kayıtlı" in cikti
    # rapor yeniden hesaplanınca (kural sürümü değişince) listeler klasörden yeniden okunur
    import json
    j = json.loads((k / depo.RAPOR_JSON).read_text(encoding="utf-8")); j["kural"] = 0
    (k / depo.RAPOR_JSON).write_text(json.dumps(j), encoding="utf-8")
    v2 = c.get("/api/rapor", params={"id": v["id"]}).json()
    assert len(v2["red_iptal"]) == 4 and v2["kural"]
