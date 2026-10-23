"""Masaüstü programının Toplu Kontrol ekranı ve API'si (sahte toplu süreçle)."""
import os
import sys
import time
from pathlib import Path

import pytest


@pytest.fixture
def istemci(tmp_path, monkeypatch):
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    monkeypatch.setenv("PYTHONPATH", os.pathsep.join([str(Path(__file__).parent), str(Path(__file__).parent.parent)]))
    from fastapi.testclient import TestClient
    from fatura_kontrol.web import app as A, toplu_api
    monkeypatch.setattr(toplu_api, "MODUL", "sahte_toplu")
    toplu_api._is.clear()
    return TestClient(A.app), toplu_api


def _bekle(c, sure=20):
    son = time.time() + sure
    while time.time() < son:
        d = c.get("/api/toplu/durum").json()
        if not d["calisiyor"]:
            return d
        time.sleep(0.3)
    raise AssertionError("iş bitmedi")


def test_program_akisi(istemci):
    c, api = istemci
    assert c.get("/api/toplu/firmalar").json()["firmalar"] == []
    assert c.get("/api/ayarlar").json()["kurulum_tamam"] is False
    assert c.post("/api/toplu/baslat", json={"ay": "2026-09"}).status_code == 400      # liste yok
    c.post("/api/toplu/firmalar/cek")
    d = _bekle(c)
    assert d["ilerleme"]["asama"] == "bitti"
    f = c.get("/api/toplu/firmalar").json()["firmalar"]
    assert [x["ad"] for x in f] == ["EXPO İNŞ", "RAKİKOM", "YENİ FİRMA"]
    # YENİ FİRMA hariç, not ekle
    c.put("/api/toplu/firmalar", json={"firmalar": [{"id": "3", "dahil": False, "not": "kapandı"}]})
    f = {x["id"]: x for x in c.get("/api/toplu/firmalar").json()["firmalar"]}
    assert f["3"]["dahil"] is False and f["3"]["not"] == "kapandı" and f["1"]["dahil"]
    # defter türü: Firmalar sekmesinden yazılır, MuhaWin'in defter_turleri ile ortak
    c.put("/api/toplu/firmalar", json={"firmalar": [{"id": "2", "dahil": True, "defter": "isletme"}, {"id": "1", "dahil": True}]})
    f = {x["id"]: x for x in c.get("/api/toplu/firmalar").json()["firmalar"]}
    assert f["2"]["defter"] == "isletme" and f["1"]["defter"] == "" and f["3"]["not"] == "kapandı"
    c.put("/api/toplu/firmalar", json={"firmalar": [{"id": "2", "dahil": True, "defter": ""}]})
    assert {x["id"]: x for x in c.get("/api/toplu/firmalar").json()["firmalar"]}["2"]["defter"] == ""

    assert c.post("/api/toplu/baslat", json={"ay": "2026-13"}).status_code == 400
    k = c.post("/api/toplu/baslat", json={"ay": "2026-09"}).json()["klasor"]
    assert k == "2026-09"
    assert c.post("/api/toplu/baslat", json={"ay": "2026-09"}).status_code == 409     # aynı anda tek iş
    d = _bekle(c)
    durumlar = {r["ad"]: r["durum"] for r in d["sonuclar"]}
    assert durumlar == {"EXPO İNŞ": "Sorun var", "RAKİKOM": "Eksik indirme"}
    rak = next(r for r in d["sonuclar"] if r["ad"] == "RAKİKOM")
    assert {a["kod"]: a["durum"] for a in rak["adimlar"]} == {
        "gib_efatura_alis": "tamam", "gib_efatura_satis": "belge_yok", "gib_ebelge_alis": "hata",
        "gib_ebelge_satis": "yapilmadi", "gib_esmm_alis": "yapilmadi", "gib_esmm_satis": "yapilmadi", "muavin": "tamam"}
    # adım tekrarı yalnız tek firma + mevcut çalışma ile
    assert c.post("/api/toplu/baslat", json={"ay": "2026-09", "idler": ["2"], "adim": "gib_ebelge_alis"}).status_code == 400
    assert c.post("/api/toplu/baslat", json={"ay": "2026-09", "idler": ["2"], "klasor": k, "adim": "xx"}).status_code == 400
    # yalnız hatalı adımı aynı çalışmada yeniden dene → diğer sonuç korunur
    c.post("/api/toplu/baslat", json={"ay": "2026-09", "idler": ["2"], "klasor": k, "adim": "gib_ebelge_alis"})
    d = _bekle(c)
    rak = next(r for r in d["sonuclar"] if r["ad"] == "RAKİKOM")
    assert {a["kod"]: a["durum"] for a in rak["adimlar"]}["gib_ebelge_alis"] == "tamam"
    kdv2 = c.get(f"/api/toplu/calisma/{k}/kdv2").json()["firmalar"]
    assert [f["ad"] for f in kdv2] == ["EXPO İNŞ"] and kdv2[0]["toplam"] == 560 and len(kdv2[0]["satirlar"]) == 1
    assert c.get(f"/api/toplu/calisma/{k}/kdv2.xlsx").status_code == 200
    durumlar = {r["ad"]: r["durum"] for r in d["sonuclar"]}
    assert durumlar == {"EXPO İNŞ": "Sorun var", "RAKİKOM": "Sorun var"}
    from openpyxl import load_workbook
    wb = load_workbook(Path(api._toplu_kok() / k / "toplu_ozet.xlsx"))
    assert "KDV2 Raporu" in wb.sheetnames
    g = c.get("/api/toplu/gecmis").json()
    assert g[0]["klasor"] == k and g[0]["firma"] == 2 and g[0]["sorunlu"] == 2 and g[0]["ozet_var"]
    assert c.get(f"/api/toplu/calisma/{k}/ozet.xlsx").status_code == 200
    assert c.get("/api/toplu/calisma/..%2F..%2Fx").status_code == 404
    # Sonuçlar tablosunda seçilenlerin sonucunu sil (isteğe bağlı: dönem klasörü de)
    from datetime import date as _d
    from fatura_kontrol import depo
    fk = depo.donem_klasoru("RAKİKOM", _d(2026, 9, 1)); fk.mkdir(parents=True, exist_ok=True); (fk / "x.txt").write_text("x")
    assert c.post(f"/api/toplu/calisma/{k}/sil", json={"anahtarlar": []}).status_code == 400
    assert c.post(f"/api/toplu/calisma/{k}/sil", json={"anahtarlar": ["yok"]}).status_code == 404
    j = c.post(f"/api/toplu/calisma/{k}/sil", json={"anahtarlar": ["2"], "belgeler": True}).json()
    assert j == {"silinen": 1, "klasor_silinen": 1, "kalan": 1} and not fk.exists()
    assert [r["ad"] for r in c.get(f"/api/toplu/calisma/{k}").json()["sonuclar"]] == ["EXPO İNŞ"]
    assert c.get("/api/toplu/gecmis").json()[0]["firma"] == 1
    assert c.post(f"/api/toplu/calisma/{k}/sil", json={"anahtarlar": ["1"]}).json()["klasor_silinen"] == 0


def test_program_ekrani(istemci, tmp_path):
    """Ekran tarayıcıda hatasız açılır, iş başlatılır, canlı ilerleme ve sonuçlar görünür."""
    pytest.importorskip("playwright")
    import threading
    import uvicorn
    from playwright.sync_api import sync_playwright
    c, api = istemci
    from fatura_kontrol.web.app import app
    srv = uvicorn.Server(uvicorn.Config(app, host="127.0.0.1", port=8799, log_level="error"))
    t = threading.Thread(target=srv.run, daemon=True); t.start()
    time.sleep(1.5)
    hatalar = []
    try:
        with sync_playwright() as p:
            b = p.chromium.launch()
            s = b.new_page(viewport={"width": 1400, "height": 900})
            s.on("pageerror", lambda e: hatalar.append(str(e)))
            s.goto("http://127.0.0.1:8799/toplu")
            s.wait_for_selector("#hosgeldin >> text=Hoş geldiniz")      # ilk açılışta Ayarlar
            s.fill("[name=ad_soyad]", "Mehmet Test"); s.fill("[name=buro_unvan]", "Test SMMM Bürosu")
            s.fill("[name=luca_uye_no]", "12345"); s.fill("[name=luca_kullanici]", "deneme")
            s.fill("[name=luca_parola]", "GizliParola1"); s.check("[name=otomatik_giris]")
            s.click("#ayarForm button[type=submit]")
            s.wait_for_selector("text=Parola kayıtlı.")
            assert "GizliParola1" not in s.content()                       # parola sayfaya geri gelmez
            assert "Test SMMM Bürosu" in s.inner_text("#kullaniciAdi")
            s.goto("http://127.0.0.1:8799/toplu")
            s.wait_for_selector("text=Liste boş")                       # liste yokken Firmalar sekmesi açılır
            s.click("#cekBtn")
            s.wait_for_selector("#firmaTablo >> text=RAKİKOM", timeout=15000)
            s.locator("#firmaTablo tr", has_text="YENİ FİRMA").locator(".dahil").uncheck()
            assert s.is_visible("#kaydedilmedi")
            s.fill("#ay", "2026-09")
            s.click("#baslatBtn")
            s.wait_for_selector("text=Luca'ya giriş yapın.", timeout=10000)
            s.wait_for_selector("#sonucTablo >> text=Eksik indirme", timeout=20000)
            s.wait_for_selector("text=Tamamlandı.", timeout=20000)
            assert s.is_visible("#ozetBtn") and s.is_visible("#hatalilarBtn")
            s.wait_for_selector("#islemOzeti .tekrarBtn")
            s.screenshot(path=os.environ.get("FK_EKRAN", str(tmp_path / "ekran.png")))
            s.click("#islemOzeti .tekrarBtn[data-adim=gib_ebelge_alis]")
            s.wait_for_function("document.querySelector('#hatalilarBtn').hidden", timeout=20000)
            s.click("#kdv2Btn")
            s.wait_for_selector("#kdv2Icerik >> text=NAKLİYE LTD")
            metin = s.locator("#kdv2Icerik").inner_text()
            assert "MÜŞTERİ" not in metin and "SONRAKİ AY" not in metin      # yalnız alışlar, yalnız o ay
            s.screenshot(path=os.environ.get("FK_EKRAN2", str(tmp_path / "ekran2.png")))
            s.click("#kdv2Icerik .raporBtn")                              # rapor aynı ekranda açılır
            assert s.is_visible("#raporPencere") and "gomulu=1" in s.get_attribute("#raporCerceve", "src")
            assert s.url.endswith("/toplu")
            s.click("#raporKapat")
            assert not s.is_visible("#raporPencere")
            s.click("text=Aylık geçmiş")
            s.wait_for_selector("#gecmisTablo >> text=Eylül 2026")
            b.close()
    finally:
        srv.should_exit = True
    assert not hatalar, hatalar
    from fatura_kontrol.luca import toplu as T
    assert {f["id"]: f["dahil"] for f in T.liste_oku()}["3"] is False
    from fatura_kontrol import ayarlar
    ham = (tmp_path / "veri" / "ayarlar.json").read_text(encoding="utf-8")
    assert "GizliParola1" not in ham
    assert ayarlar.luca_girisi() == {"uye_no": "12345", "kullanici": "deneme", "parola": "GizliParola1"}


def test_kayit_bildirimi(tmp_path, monkeypatch):
    """Kullanıcı bilgileri kaydedilince form servisine gönderilir; Luca bilgileri gönderilmez;
    değişmeyen bilgi tekrar gönderilmez; başarısız gönderim sonra yeniden denenir."""
    import json
    import threading
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
    gelen = []

    class H(BaseHTTPRequestHandler):
        def do_POST(self):
            gelen.append(json.loads(self.rfile.read(int(self.headers["Content-Length"])).decode("utf-8")))
            self.send_response(200); self.send_header("Content-Type", "application/json"); self.end_headers()
            self.wfile.write(b'{"success":"true"}')

        def log_message(self, *a):
            pass
    srv = ThreadingHTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fatura_kontrol import ayarlar, bildirim
    monkeypatch.setenv("FK_BILDIRIM_URL", "http://127.0.0.1:1/yok")           # önce internet yok
    ayarlar.kaydet({"ad_soyad": "Ali Veli", "buro_unvan": "Veli SMMM", "eposta": "ali@ornek.com",
                    "luca_uye_no": "999", "luca_kullanici": "gizlikul", "luca_parola": "GizliP4rola"})
    assert bildirim.gonder()["hata"]
    monkeypatch.setenv("FK_BILDIRIM_URL", f"http://127.0.0.1:{srv.server_port}/f")
    assert bildirim.gonder()["gonderildi"] and len(gelen) == 1                  # yeniden denendi
    metin = json.dumps(gelen[0], ensure_ascii=False)
    assert "Veli SMMM" in metin and "ali@ornek.com" in metin
    assert "999" not in metin and "gizlikul" not in metin and "GizliP4rola" not in metin
    bildirim.gonder()
    assert len(gelen) == 1                                                      # değişmedi → gönderilmez
    ayarlar.kaydet({"telefon": "0242 000 00 00"})
    bildirim.gonder()
    assert len(gelen) == 2 and gelen[1]["Durum"] == "bilgi güncellemesi"
    srv.shutdown()


def test_cikti_raporu_ve_eski_raporun_yenilenmesi(istemci, tmp_path):
    """Eski kurallarla yazılmış rapor açılınca klasördeki fatura + muavinle yeniden hesaplanır;
    çıktı raporu (özet + muavin dökümü + KDV beyannamesi örneği) üretilir ve klasöre kaydedilir."""
    import json, zipfile
    from datetime import date
    sys.path.insert(0, str(Path(__file__).parent))
    from test_luca_otomasyon import muavin_yaz, ubl, FIRMA_VKN
    from fatura_kontrol import depo
    c, _ = istemci
    k = depo.donem_klasoru("EXPO İNŞ", date(2026, 9, 1)); k.mkdir(parents=True)
    muavin_yaz(k / "muavin.xlsx")
    with zipfile.ZipFile(k / "faturalar.zip", "w") as z:
        z.writestr("gib_efatura_alis/AAA2026000000001.xml",
                   ubl("AAA2026000000001", "2026-09-05", "1", FIRMA_VKN, "1000.00", "200.00", 20))
    eski = {"firma": "EXPO", "vkn": FIRMA_VKN, "donem": {"Alış": ["2026-09-01", "2026-09-30"]}, "uyarilar": [],
            "durumlar": [], "sonuclar": [], "fatura_sayisi": 0, "id": "EXPO İNŞ/2026/09"}
    (k / depo.RAPOR_JSON).write_text(json.dumps(eski), encoding="utf-8")
    v = c.get("/api/rapor", params={"id": "EXPO İNŞ/2026/09"}).json()
    assert v["fatura_sayisi"] == 1 and v["kural"] and v["sonuclar"][0]["durum"] == "Uyumlu"
    r = c.get("/api/rapor/cikti", params={"id": "EXPO İNŞ/2026/09"})
    assert r.status_code == 200 and "KDV beyannamesi örneği" in r.text and "AAA2026000000001" in r.text
    assert (k / "kdv_kontrol_raporu.html").exists()
    assert c.get("/api/rapor/cikti", params={"id": "YOK/2026/09"}).status_code == 404
    c.get("/api/rapor/cikti", params={"id": "EXPO İNŞ/2026/09", "cevap": "hayir"})
    assert json.loads((k / "iade_talebi.json").read_text(encoding="utf-8"))["cevap"] == "hayir"


def test_tek_firma_toplu_ile_ayni(istemci, tmp_path):
    """Tek firma (dosyadan) kontrol toplu kontrolle aynı yolu izler: dosyalar dönem klasörüne arşivlenir,
    ihraç kayıtlı / istisna gibi fatura tipleri aynı sınıflanır, çıktı raporunda muavin dökümü ve KDV beyannamesi vardır."""
    import io, json, zipfile
    sys.path.insert(0, str(Path(__file__).parent))
    from test_luca_otomasyon import muavin_yaz, ubl, FIRMA_VKN
    from fatura_kontrol import depo
    c, _ = istemci
    m = tmp_path / "muavin.xlsx"; muavin_yaz(m)
    ihrac = ubl("EX02026000000002", "2026-09-12", FIRMA_VKN, "9", "5000.00", "1000.00", 20).replace(
        b"<cbc:InvoiceTypeCode>SATIS</cbc:InvoiceTypeCode>", b"<cbc:InvoiceTypeCode>IHRACKAYITLI</cbc:InvoiceTypeCode>")
    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w") as z:
        z.writestr("a.xml", ubl("AAA2026000000001", "2026-09-05", "1", FIRMA_VKN, "1000.00", "200.00", 20))
        z.writestr("b.xml", ubl("EX02026000000001", "2026-09-10", FIRMA_VKN, "9", "2000.00", "400.00", 20))
    assert c.get("/api/muavin/programlar").json()["programlar"][0]["kod"] == "luca"
    bad = c.post("/api/kontrol", data={"program": "zirve"}, files=[("faturalar", ("p.zip", b"x", "application/zip")),
                                                                 ("muavinler", ("m.xlsx", b"x", "application/octet-stream"))])
    assert bad.status_code == 400 and "Bilinmeyen muavin programı" in bad.text
    r = c.post("/api/kontrol", data={"bas": "2026-09-01", "bit": "2026-09-30", "program": "luca"},
               files=[("faturalar", ("paket.zip", zb.getvalue(), "application/zip")),
                      ("faturalar", ("ihrac.xml", ihrac, "application/xml")),
                      ("muavinler", ("muavin.xlsx", m.read_bytes(), "application/octet-stream"))])
    assert r.status_code == 200, r.text
    v = r.json()
    k = depo.rapor_klasoru(v["id"])
    assert (k / "faturalar.zip").exists() and (k / "muavin.xlsx").exists()
    assert json.loads((k / "muavin_program.json").read_text())["program"] == "luca"
    assert not list(k.glob("gib_*.zip"))                     # ham paketler arşive alındı
    tipler = {f["belge_no"]: f["fatura_tipi"] for f in v["faturalar"]}
    assert "İhraç kayıtlı" in tipler["EX02026000000002"]
    d = {s["belge_no"]: s["durum"] for s in v["sonuclar"]}
    assert "ihra" in d["EX02026000000002"].lower() or "istisna" in d["EX02026000000002"].lower()   # KDV'siz ihraç sınıfı
    cikti = c.get("/api/rapor/cikti", params={"id": v["id"]}).text
    assert "KDV beyannamesi örneği" in cikti and "muavin" in cikti.lower()
    g = c.get("/api/fatura/goruntu", params={"id": v["id"], "belge": "AAA2026000000001"})
    assert g.status_code == 200


def test_tek_firma_faturalari_lucadan(istemci, tmp_path):
    """Tek firma › "Faturaları Luca'dan Al": toplu otomasyon yalnız e-belgeleri indirir (muavin dokunulmaz),
    kontrol yüklenen dosya olmadan firmanın dönem arşivindeki faturalarla yapılır."""
    import json
    sys.path.insert(0, str(Path(__file__).parent))
    from test_luca_otomasyon import muavin_yaz
    from fatura_kontrol import depo
    from fatura_kontrol.luca import toplu as T
    c, api = istemci
    T.liste_yaz([{"ad": "EXPO İNŞ", "id": "1"}, {"ad": "RAKİKOM", "id": "2"}])
    assert c.post("/api/toplu/baslat", json={"ay": "2026-09", "yalniz_fatura": True}).status_code == 400          # firma yok
    assert c.post("/api/toplu/baslat", json={"ay": "2026-09", "idler": ["9"], "yalniz_fatura": True}).status_code == 400
    assert c.get("/api/kontrol/luca-faturalari", params={"firma_id": "1", "bas": "2026-09-01"}).json()["adet"] == 0
    r = c.post("/api/toplu/baslat", json={"ay": "2026-09", "idler": ["1"], "yalniz_fatura": True})
    assert r.json()["klasor"] == "_tek/2026-09"
    d = _bekle(c)
    assert d["cikis_kodu"] == 0, d["gunluk"]
    assert "--yalniz-fatura" in " ".join(api._is["proc"].args)
    assert c.get("/api/toplu/gecmis").json() == []                     # toplu kontrol geçmişine girmez
    j = c.get("/api/kontrol/luca-faturalari", params={"firma_id": "1", "bas": "2026-09-01"}).json()
    assert j["adet"] == 2 and j["turler"] == {"gib_efatura_alis": 1, "gib_efatura_satis": 1}
    assert [x["kod"] for x in j["eksik"]] == ["gib_efatura_satis"]          # indirilemeyen tür → "Luca'dan tekrar dene"
    assert c.post("/api/toplu/baslat", json={"ay": "2026-09", "idler": ["1"], "yalniz_fatura": True, "adim": "muavin"}).status_code == 400
    c.post("/api/toplu/baslat", json={"ay": "2026-09", "idler": ["1"], "yalniz_fatura": True, "adim": "gib_efatura_satis"})
    d = _bekle(c)
    assert d["cikis_kodu"] == 0 and "gib_efatura_satis" in api._is["proc"].args
    assert c.get("/api/kontrol/luca-faturalari", params={"firma_id": "1", "bas": "2026-09-01"}).json()["eksik"] == []
    m = tmp_path / "muavin.xlsx"; muavin_yaz(m)
    yok = c.post("/api/kontrol", data={"bas": "2026-10-01", "bit": "2026-10-31", "program": "luca", "firma_id": "1", "fatura_kaynak": "luca"},
                 files=[("muavinler", ("muavin.xlsx", m.read_bytes(), "application/octet-stream"))])
    assert yok.status_code == 400 and "Luca'dan alınmış fatura yok" in yok.text
    r = c.post("/api/kontrol", data={"bas": "2026-09-01", "bit": "2026-09-30", "program": "luca", "firma_id": "1", "firma_ad": "EXPO İNŞ", "fatura_kaynak": "luca"},
               files=[("muavinler", ("muavin.xlsx", m.read_bytes(), "application/octet-stream"))])
    assert r.status_code == 200, r.text
    v = r.json()
    assert v["fatura_sayisi"] == 2 and v["id"] == "EXPO İNŞ/2026/09"
    k = depo.rapor_klasoru(v["id"])
    assert (k / "faturalar.zip").exists() and (k / "muavin.xlsx").exists()
    assert c.get("/api/kontrol/luca-faturalari", params={"firma_id": "1", "bas": "2026-09-01"}).json()["adet"] == 2   # arşiv korunur
    assert c.get("/api/fatura/goruntu", params={"id": v["id"], "belge": "AAA2026000000001"}).status_code == 200


def test_kayit_adaylari(istemci, tmp_path):
    """Fatura → Luca kayıt adayları: kayıtsız faturalar işaretlenir, geçmiş kayıtlardan hesap önerisi öğrenilir."""
    import io, zipfile
    sys.path.insert(0, str(Path(__file__).parent))
    from test_luca_otomasyon import muavin_yaz, ubl, FIRMA_VKN
    c, _ = istemci
    m = tmp_path / "muavin.xlsx"; muavin_yaz(m)
    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w") as z:
        z.writestr("a.xml", ubl("AAA2026000000001", "2026-09-05", "1", FIRMA_VKN, "1000.00", "200.00", 20))
        z.writestr("y.xml", ubl("YYY2026000000009", "2026-09-20", "1", FIRMA_VKN, "500.00", "100.00", 20))
    r = c.post("/api/kontrol", data={"bas": "2026-09-01", "bit": "2026-09-30"},
               files=[("faturalar", ("p.zip", zb.getvalue(), "application/zip")),
                      ("muavinler", ("muavin.xlsx", m.read_bytes(), "application/octet-stream"))])
    v = r.json()
    a = c.get("/api/kayit/adaylar", params={"id": v["id"]}).json()
    f = {x["belge_no"]: x for x in a["faturalar"]}
    assert f["AAA2026000000001"]["kayitli"] and not f["YYY2026000000009"]["kayitli"]
    assert f["YYY2026000000009"]["oranlar"] == [{"oran": "20", "matrah": 500.0, "kdv": 100.0}]
    assert a["kdv_hesaplari"]["ALIS"]["20"] == "191.10.20"
    assert a["kdv_hesaplari"]["SATIS"]["20"] == "391.10.20"
    assert c.get("/api/kayit/adaylar", params={"id": "YOK/2026/09"}).status_code == 404


def test_mikro_190_devreden_ve_iade_kdv(tmp_path):
    """Mikro muavininde 190'ın DEVREDEN satırı dönem başı bakiyesi olarak okunur; 'Satışlardan İade KDV' 102 indirimidir."""
    from datetime import datetime, date
    from decimal import Decimal as D
    import openpyxl
    from fatura_kontrol.mikro_muavin import mikro_oku
    from fatura_kontrol.cikti import indirim_turu, kdv_beyani
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["HESAP KODU", "Kayıt no", "Tarih", "Sıra", "Y.No", "Fiş tipi", "Belge türü", "Evrak no", "Cari ismi", "Entegrasyon kaynağı", "Açıklama", "TL Borç", "TL Alacak"])
    ws.append(["190.01 Devreden Kdv Hesap hareket dökümü  (01.09.2026-30.09.2026)"])
    ws.append(["", "", datetime(2026, 9, 1), None, None, "", "", "", "", "", " D E V R E D E N", 5000, 1000])
    ws.append(["391.01 Hesaplanan KDV (%20) Hesap hareket dökümü  (01.09.2026-30.09.2026)"])
    ws.append(["", "K1", datetime(2026, 9, 5), 1, 10, "Mahsup", "Sat.fat", "1", "ALICI", "", "Sat.fat : AAS2026000000001/ALICI", None, 2000])
    ws.append(["191.02.20 Satışlardan İade KDV Hesabı % 20 Hesap hareket dökümü  (01.09.2026-30.09.2026)"])
    ws.append(["", "K2", datetime(2026, 9, 6), 1, 11, "Mahsup", "Al.fat.", "M012026-29", "X", "", "Al.fat. : M012026000000029/X", 100, None])
    y = tmp_path / "m.xlsx"; wb.save(y)
    s, _ = mikro_oku(y)
    dv = [x for x in s if x.hesap_kodu.startswith("190")]
    assert len(dv) == 1 and dv[0].fis_tarihi == date(2026, 8, 31) and dv[0].borc - dv[0].alacak == 4000
    assert next(x for x in s if x.hesap_kodu.startswith("191")).belge_no == "M012026000000029"
    assert indirim_turu("191.02.20", "Satışlardan İade KDV Hesabı % 20").startswith("102")
    assert indirim_turu("191.02.01", "Sorumlu Sıfatıyla Ödenen KDV").startswith("109")
    b = kdv_beyani(s, date(2026, 9, 1), date(2026, 9, 30))
    assert b.devreden == D("4000") and b.hesaplanan == D("2000")
