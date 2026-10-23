"""Luca otomasyonunun sahte Luca sitesine karşı uçtan uca testi."""
import json
import os
from datetime import date
from pathlib import Path

import pytest

pytest.importorskip("playwright")

FIRMA_VKN = "3810463613"


def ubl(no, tarih, satici, alici, matrah, kdv, oran):
    return f"""<?xml version="1.0" encoding="UTF-8"?>
<Invoice xmlns="urn:oasis:names:specification:ubl:schema:xsd:Invoice-2"
 xmlns:cac="urn:oasis:names:specification:ubl:schema:xsd:CommonAggregateComponents-2"
 xmlns:cbc="urn:oasis:names:specification:ubl:schema:xsd:CommonBasicComponents-2">
<cbc:ProfileID>TEMELFATURA</cbc:ProfileID><cbc:ID>{no}</cbc:ID><cbc:UUID>{no}-uuid</cbc:UUID>
<cbc:IssueDate>{tarih}</cbc:IssueDate><cbc:InvoiceTypeCode>SATIS</cbc:InvoiceTypeCode><cbc:DocumentCurrencyCode>TRY</cbc:DocumentCurrencyCode>
<cac:AccountingSupplierParty><cac:Party><cac:PartyIdentification><cbc:ID schemeID="VKN">{satici}</cbc:ID></cac:PartyIdentification>
<cac:PartyName><cbc:Name>SATICI {satici}</cbc:Name></cac:PartyName></cac:Party></cac:AccountingSupplierParty>
<cac:AccountingCustomerParty><cac:Party><cac:PartyIdentification><cbc:ID schemeID="VKN">{alici}</cbc:ID></cac:PartyIdentification>
<cac:PartyName><cbc:Name>ALICI {alici}</cbc:Name></cac:PartyName></cac:Party></cac:AccountingCustomerParty>
<cac:TaxTotal><cbc:TaxAmount currencyID="TRY">{kdv}</cbc:TaxAmount><cac:TaxSubtotal>
<cbc:TaxableAmount currencyID="TRY">{matrah}</cbc:TaxableAmount><cbc:TaxAmount currencyID="TRY">{kdv}</cbc:TaxAmount><cbc:Percent>{oran}</cbc:Percent>
<cac:TaxCategory><cac:TaxScheme><cbc:TaxTypeCode>0015</cbc:TaxTypeCode></cac:TaxScheme></cac:TaxCategory></cac:TaxSubtotal></cac:TaxTotal>
<cac:LegalMonetaryTotal><cbc:TaxExclusiveAmount currencyID="TRY">{matrah}</cbc:TaxExclusiveAmount>
<cbc:TaxInclusiveAmount currencyID="TRY">{float(matrah) + float(kdv):.2f}</cbc:TaxInclusiveAmount>
<cbc:PayableAmount currencyID="TRY">{float(matrah) + float(kdv):.2f}</cbc:PayableAmount></cac:LegalMonetaryTotal>
</Invoice>""".encode()


def muavin_yaz(yol: Path):
    from datetime import datetime
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    for r in [("MUAVİN DEFTER",), ("EXPO İNŞAAT SANAYİ VE TİCARET LİMİTED ŞİRKETİ",), ("Dönem :", "01/01/2026-31/12/2026"),
              ("Tarih Aralığı :", "01/01/2026-31/12/2026"), ("191.10.20 %20 İndirilecek KDV",),
              ("TARİH", "TİP", "FİŞ NO", "AÇIKLAMA", "BORÇ", "ALACAK", "BAKİYE", "B/A"),
              (datetime(2026, 9, 30), "Mahsup", "1", "05/09/2026-AAA2026000000001-SATICI 1", 200, None, 0, 0),
              (datetime(2026, 9, 30), "Mahsup", "1", "06/09/2026-AAA2026000000002-SATICI 2", 150, None, 0, 0),
              ("391.10.20 %20 Hesaplanan KDV",),
              ("TARİH", "TİP", "FİŞ NO", "AÇIKLAMA", "BORÇ", "ALACAK", "BAKİYE", "B/A"),
              (datetime(2026, 9, 30), "Mahsup", "2", "10/09/2026-EX02026000000001-ALICI 9", None, 400, 0, 0)]:
        ws.append(r)
    wb.save(yol)


@pytest.fixture
def sahte_luca(tmp_path):
    from luca_sahte import site_kur
    m = tmp_path / "muavin.xlsx"; muavin_yaz(m)
    xmller = {
        "gib_efatura_alis_1.zip": [ubl("AAA2026000000001", "2026-09-05", "1", FIRMA_VKN, "1000.00", "200.00", 20)],
        "gib_efatura_alis_2.zip": [ubl("AAA2026000000002", "2026-09-06", "2", FIRMA_VKN, "1000.00", "200.00", 20)],
        "gib_efatura_alis_3.zip": [ubl("AAA2026000000003", "2026-09-07", "3", FIRMA_VKN, "500.00", "100.00", 20)],
        "gib_efatura_satis_1.zip": [ubl("EX02026000000001", "2026-09-10", FIRMA_VKN, "9", "2000.00", "400.00", 20)],
    }
    veri = {"gib_efatura_alis": 5, "gib_efatura_satis": 1, "gib_ebelge_alis": 0, "gib_ebelge_satis": 0,
            "gib_esmm_alis": 0, "gib_esmm_satis": 0}
    srv, url = site_kur(tmp_path / "site", xmller, m, veri)
    yield url
    srv.shutdown()


def test_uctan_uca(sahte_luca, tmp_path, monkeypatch):
    monkeypatch.setenv("FK_HEADLESS", "1")
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fatura_kontrol.luca import otomasyon as O
    monkeypatch.setattr(O, "BASLANGIC", sahte_luca)
    monkeypatch.setattr(O, "GETIR_BEKLEME", 25)

    log = {}

    def giris(ctx, sayfa):
        sayfa.fill("#uye", "1212207"); sayfa.fill("#sifre", "x")
        with sayfa.expect_popup():
            sayfa.click("#giris")

    def log_al(ctx):
        for p in ctx.pages:
            if "luca.do" in p.url:
                log["v"] = p.evaluate("JSON.parse(sessionStorage.getItem('log') || '[]')")

    orijinal_muavin = O.muavin_indir

    def muavin_ve_log(L, kok, uye, s):
        orijinal_muavin(L, kok, uye, s)
        log["v"] = L.evaluate("JSON.parse(sessionStorage.getItem('log') || '[]')")
    monkeypatch.setattr(O, "muavin_indir", muavin_ve_log)

    s = O.indir("expo inş", date(2026, 9, 1), date(2026, 9, 30), list(O.TURLER), tmp_path / "indir",
                profil=tmp_path / "profil", ctx_hazirla=giris)
    adimlar = {a["adim"]: a for a in s.adimlar}
    print(json.dumps(s.adimlar, ensure_ascii=False, indent=1))

    assert s.firma == "EXPO İNŞ" and s.vkn == FIRMA_VKN
    assert [p.name for p in s.zipler] == ["gib_efatura_alis_1.zip", "gib_efatura_alis_2.zip",
                                          "gib_efatura_alis_3.zip", "gib_efatura_satis_1.zip"]
    assert s.muavin and s.muavin.exists()
    assert adimlar["e-Arşiv Alış: GİB sorgusu"]["durum"] == "tamam"   # sorgu sırasındaki pencere geçildi
    assert adimlar["e-Arşiv Alış: indirme"]["durum"] == "atlandı"
    assert adimlar["e-Arşiv Satış"]["durum"] == "hata"          # bitmeyen sorgu, akış devam etti
    assert adimlar["e-SMM Alış: indirme"]["durum"] == "atlandı"  # 'bulunamadı' penceresi Tamam ile geçildi
    assert adimlar["e-SMM Satış: indirme"]["durum"] == "atlandı"
    olaylar = log["v"]
    assert not any(k == "YASAK" for k, _ in olaylar)
    getirler = [v for k, v in olaylar if k == "getir"]
    assert ["gib_efatura_alis", "01/09/2026", "30/09/2026", 110198728] in getirler
    bugun = date.today()
    if bugun.replace(day=1) > date(2026, 9, 30):                    # içinde bulunulan ay da sorgulandı
        assert ["gib_efatura_alis", bugun.replace(day=1).strftime("%d/%m/%Y"), bugun.strftime("%d/%m/%Y"),
                110198728] in getirler
    assert not any(k == "YASAK" for k, _ in olaylar)                # 'İptal' düğmesine basılmadı
    assert all(v[3] == 110198728 for v in getirler)                 # hep doğru firmada

    veri = O.kontrol_et(s, date(2026, 9, 1), date(2026, 9, 30))
    sayac = {d["ad"]: d["adet"] for d in veri["durumlar"]}
    assert sayac["Uyumlu"] == 2 and sayac["Tutar farkı"] == 1 and sayac["Muavinde yok"] == 1
    assert any("e-Arşiv Satış" in u for u in veri["uyarilar"])
    from fatura_kontrol import depo
    assert veri["id"].endswith("/2026/09")
    assert (depo.rapor_klasoru(veri["id"]) / "kontrol_raporu.xlsx").exists()


def test_sorgu_araliklari():
    from fatura_kontrol.luca.otomasyon import sorgu_araliklari
    assert sorgu_araliklari(date(2026, 9, 1), date(2026, 9, 30), date(2026, 10, 16)) == [
        (date(2026, 9, 1), date(2026, 9, 30)), (date(2026, 10, 1), date(2026, 10, 16))]
    assert sorgu_araliklari(date(2026, 10, 1), date(2026, 10, 31), date(2026, 10, 16)) == [
        (date(2026, 10, 1), date(2026, 10, 16))]
    assert sorgu_araliklari(date(2026, 7, 1), date(2026, 7, 31), date(2026, 10, 16))[-1] == (
        date(2026, 10, 1), date(2026, 10, 16))


def test_yasakli_dugmeye_basilmaz(sahte_luca, tmp_path, monkeypatch):
    monkeypatch.setenv("FK_HEADLESS", "1")
    from playwright.sync_api import sync_playwright
    from fatura_kontrol.luca.otomasyon import OtomasyonHatasi, guvenli_tikla
    with sync_playwright() as p:
        b = p.chromium.launch()
        pg = b.new_page()
        pg.goto(sahte_luca.replace("index.html", "Luca/gib530.do?tur=gib_efatura_alis"))
        for sec in ("#sil-btn", "#muhasebelestir-btn"):
            with pytest.raises(OtomasyonHatasi, match="GÜVENLİK"):
                guvenli_tikla(pg.locator(sec), sec)
        assert pg.evaluate("sessionStorage.getItem('log')") is None
        b.close()


def test_firma_listesi_dosyasi(tmp_path):
    from fatura_kontrol.luca import toplu as T
    yol = tmp_path / "firmalar.xlsx"
    T.liste_yaz([{"id": "1", "ad": "A FİRMA"}, {"id": "2", "ad": "B FİRMA"}], yol)
    from openpyxl import load_workbook
    wb = load_workbook(yol); wb.active["D3"] = "H"; wb.active["E3"] = "kapanmış"; wb.save(yol)
    T.liste_yaz([{"id": "1", "ad": "A FİRMA"}, {"id": "2", "ad": "B FİRMA"}, {"id": "3", "ad": "C YENİ"}], yol)
    f = {x["id"]: x for x in T.liste_oku(yol)}
    assert f["1"]["dahil"] and not f["2"]["dahil"] and f["2"]["not"] == "kapanmış" and f["3"]["dahil"]


def test_toplu_calistirma(sahte_luca, tmp_path, monkeypatch):
    monkeypatch.setenv("FK_HEADLESS", "1")
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fatura_kontrol.luca import otomasyon as O
    from fatura_kontrol.luca import toplu as T
    monkeypatch.setattr(O, "BASLANGIC", sahte_luca)
    monkeypatch.setattr(O, "GETIR_BEKLEME", 25)

    def giris(ctx, sayfa):
        with sayfa.expect_popup():
            sayfa.click("#giris")

    with O.oturum(tmp_path / "profil", giris) as (ctx, L):
        liste = O.firma_listesi(L)
    assert [f["ad"] for f in liste] == ["BANKOĞLU M", "EXPO İNŞ", "EXPO İNŞ ESKİ"]   # boş seçenek yok
    T.liste_yaz(liste)
    firmalar = [f for f in T.liste_oku() if f["id"] in ("110198728", "110304730")]
    turler = ["gib_efatura_alis", "gib_efatura_satis"]
    sonuclar, ozet = T.toplu_calistir(date(2026, 9, 1), date(2026, 9, 30), turler, tmp_path / "toplu",
                                      firmalar, profil=tmp_path / "profil2", ctx_hazirla=giris)
    d = {r["ad"]: r for r in sonuclar}
    assert d["EXPO İNŞ"]["vkn"] == FIRMA_VKN and d["EXPO İNŞ"]["durum"] == "Sorun var"
    assert d["EXPO İNŞ"]["sorunlar"]["Tutar farkı"] == 1 and d["EXPO İNŞ"]["rapor"]
    assert d["BANKOĞLU M"]["vkn"] == "1111111111" and d["BANKOĞLU M"]["durum"] == "Belge/muavin yok"
    assert ozet.exists() and (tmp_path / "toplu" / d["EXPO İNŞ"]["rapor"]).exists()
    from openpyxl import load_workbook
    ws = load_workbook(ozet).active
    assert [ws.cell(5, 1).value, ws.cell(6, 1).value] == ["EXPO İNŞ", "BANKOĞLU M"]   # sorunlu üstte
    adimlar = {a["kod"]: a["durum"] for a in d["EXPO İNŞ"]["adimlar"]}
    assert adimlar["gib_efatura_alis"] == "tamam" and adimlar["muavin"] == "tamam"

    # yalnız bir adımı tekrar dene: diğer adımların dosyaları ve kayıtları korunur, kontrol yeniden yapılır
    expo = [f for f in firmalar if f["ad"] == "EXPO İNŞ"]
    sonuclar2, _ = T.toplu_calistir(date(2026, 9, 1), date(2026, 9, 30), turler, tmp_path / "toplu",
                                    expo, profil=tmp_path / "profil2", ctx_hazirla=giris, adim="gib_efatura_satis")
    d2 = {r["ad"]: r for r in sonuclar2}
    assert set(d2) == {"EXPO İNŞ", "BANKOĞLU M"}
    assert d2["EXPO İNŞ"]["fatura"] == d["EXPO İNŞ"]["fatura"]
    assert d2["EXPO İNŞ"]["sorunlar"] == d["EXPO İNŞ"]["sorunlar"]
    assert {a["kod"]: a["durum"] for a in d2["EXPO İNŞ"]["adimlar"]} == adimlar
    from fatura_kontrol import depo
    k = depo.donem_klasoru("EXPO İNŞ", date(2026, 9, 1))
    assert k == tmp_path / "veri" / "firmalar" / "EXPO İNŞ" / "2026" / "09"
    assert (k / "kontrol_raporu.xlsx").exists() and (k / "muavin.xlsx").exists()
    kayit = json.loads((k / "adimlar.json").read_text(encoding="utf-8"))
    assert sum("e-Fatura Satış: GİB sorgusu" in a["adim"] for a in kayit) >= 1
    assert any(a["adim"].startswith("e-Fatura Alış") for a in kayit)

    # Tek firma › "Faturaları Luca'dan Al": yalnız e-belgeler yeniden indirilir, muavin ve muavin adımı korunur
    m0 = (k / "muavin.xlsx").stat().st_mtime_ns
    sonuclar3, _ = T.toplu_calistir(date(2026, 9, 1), date(2026, 9, 30), turler, tmp_path / "tek",
                                    expo, profil=tmp_path / "profil2", ctx_hazirla=giris, yalniz_fatura=True)
    d3 = {r["ad"]: r for r in sonuclar3}
    assert set(d3) == {"EXPO İNŞ"} and d3["EXPO İNŞ"]["fatura"] == d["EXPO İNŞ"]["fatura"]
    assert (k / "muavin.xlsx").stat().st_mtime_ns == m0 and (k / "faturalar.zip").exists() and not list(k.glob("gib_*.zip"))
    assert {a["kod"]: a["durum"] for a in d3["EXPO İNŞ"]["adimlar"]} == adimlar


def test_sahis_firmasi_tckn(tmp_path, monkeypatch):
    """Luca VKN gösterir, faturalarda TCKN yazar → kontrol faturalardaki kimlikle yapılır."""
    import zipfile
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fatura_kontrol.luca import otomasyon as O
    tckn = "24148451784"
    z = tmp_path / "a.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("1.xml", ubl("AAA2026000000001", "2026-09-05", "1234567890", tckn, "1000", "200", "20"))
        zf.writestr("2.xml", ubl("EX02026000000001", "2026-09-10", tckn, "9876543210", "2000", "400", "20"))
    m = tmp_path / "muavin.xlsx"; muavin_yaz(m)
    s = O.Sonuc(tmp_path / "c"); s.vkn = "2500367607"; s.zipler = [z]; s.muavin = m
    veri = O.kontrol_et(s, date(2026, 9, 1), date(2026, 9, 30))
    assert veri["fatura_sayisi"] == 2


def test_tevkifat_okuma_ve_adim_ozeti():
    from fatura_kontrol.fatura_okuyucu import xml_oku
    from fatura_kontrol.luca.toplu import adim_ozeti
    x = ubl("ABC2026000000001", "2026-09-12", "1234567890", FIRMA_VKN, "14000", "2800", "20").replace(
        b"<cac:LegalMonetaryTotal>", """<cac:WithholdingTaxTotal><cbc:TaxAmount currencyID="TRY">560.00</cbc:TaxAmount>
<cac:TaxSubtotal><cbc:TaxableAmount currencyID="TRY">2800.00</cbc:TaxableAmount><cbc:TaxAmount currencyID="TRY">560.00</cbc:TaxAmount>
<cbc:Percent>20</cbc:Percent><cac:TaxCategory><cac:TaxScheme><cbc:Name>YÜK TAŞIMACILIĞI HİZMETİ</cbc:Name>
<cbc:TaxTypeCode>624</cbc:TaxTypeCode></cac:TaxScheme></cac:TaxCategory></cac:TaxSubtotal></cac:WithholdingTaxTotal>
<cac:LegalMonetaryTotal>""".encode())
    f = xml_oku(x, "t.xml")
    assert str(f.tevkifat) == "560.00"
    d = f.tevkifat_detay[0]
    assert d["kod"] == "624" and d["oran"] == "2/10" and str(d["kdv"]) == "2800.00" and str(d["tutar"]) == "560.00"
    oz = {o["kod"]: o for o in adim_ozeti([
        {"adim": "e-Fatura Alış: GİB sorgusu", "durum": "tamam", "ayrinti": ""},
        {"adim": "e-Fatura Alış: sayfa 1 indirildi", "durum": "tamam", "ayrinti": "100 belge, 9 KB"},
        {"adim": "e-Fatura Alış: sayfa 2 indirildi", "durum": "tamam", "ayrinti": "8 belge, 1 KB"},
        {"adim": "e-Arşiv Alış: GİB sorgusu", "durum": "tamam", "ayrinti": ""},
        {"adim": "e-SMM Alış", "durum": "hata", "ayrinti": "zaman aşımı"}])}
    assert oz["gib_efatura_alis"]["durum"] == "tamam" and oz["gib_efatura_alis"]["ayrinti"] == "108 belge"
    assert oz["gib_ebelge_alis"]["durum"] == "yarim" and oz["gib_esmm_alis"]["durum"] == "hata"
    assert oz["muavin"]["durum"] == "yapilmadi"


def test_otomatik_giris_captcha_kullanicida(sahte_luca, tmp_path, monkeypatch):
    """Ayarlarda Luca bilgileri varsa program üye no/kullanıcı/parolayı doldurur ve GİRİŞ'e basar;
    captcha'yı DOLDURMAZ, kutuya odaklanıp kullanıcıdan ister. Sonra Mali Müşavir Paketi'ni açar."""
    monkeypatch.setenv("FK_HEADLESS", "1")
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fatura_kontrol import ayarlar
    from fatura_kontrol.luca import otomasyon as O
    ayarlar.kaydet({"luca_uye_no": "1212207", "luca_kullanici": "mehmet", "luca_parola": "Gizli*123",
                    "otomatik_giris": True})
    monkeypatch.setattr(O, "BASLANGIC", sahte_luca.replace("index.html", "giris2.html"))
    olaylar = []
    monkeypatch.setattr(O, "DINLEYICI", lambda o, a: olaylar.append(o))
    gorulen = {}

    def kullanici(ctx, sayfa):                     # kullanıcı yalnız doğrulama kodunu yazar
        p = next(p for p in ctx.pages if "giris.erp" in p.url)
        kutu = p.locator("#guvenlikKodu")
        kutu.wait_for(state="visible", timeout=10000)
        gorulen["kod_bos"] = kutu.input_value() == ""
        gorulen["odak"] = p.evaluate("document.activeElement && document.activeElement.id")
        gorulen["giris"] = json.loads(p.evaluate("localStorage.getItem('giris')"))
        kutu.fill("7K3P"); kutu.press("Enter")

    with O.oturum(tmp_path / "profil", kullanici) as (ctx, L):
        firmalar = O.firma_listesi(L)
    assert gorulen["giris"] == ["1212207", "mehmet", "Gizli*123"]
    assert gorulen["kod_bos"] and gorulen["odak"] == "guvenlikKodu"
    assert "captcha" in olaylar and "giris_tamam" in olaylar
    assert any(f["ad"] == "EXPO İNŞ" for f in firmalar)


def test_otomatik_giris_kapaliysa_dokunulmaz(sahte_luca, tmp_path, monkeypatch):
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fatura_kontrol import ayarlar
    ayarlar.kaydet({"luca_uye_no": "1", "luca_kullanici": "a", "luca_parola": "b", "otomatik_giris": False})
    assert ayarlar.luca_girisi() is None
    assert ayarlar.oku()["luca_parola_kayitli"] and "luca_parola" not in ayarlar.oku()


def test_depo_duzeni(tmp_path, monkeypatch):
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fatura_kontrol import depo
    k = depo.donem_klasoru('EXPO İNŞ: "A/B"', date(2026, 9, 1))
    assert k == tmp_path / "veri" / "firmalar" / "EXPO İNŞ A B" / "2026" / "09"
    k.mkdir(parents=True); (k / "eski.zip").write_bytes(b"x" * 10)
    depo.donem_temizle(k)
    assert k.exists() and not list(k.iterdir())                       # önceki indirmeler silindi
    assert depo.rapor_id(k) == "EXPO İNŞ A B/2026/09"
    assert depo.rapor_klasoru("EXPO İNŞ A B/2026/09") == k.resolve()
    for kotu in ("../x/2026/09", "a/../../2026/09", "a/2026", "a/26/09", "a/2026/09/x"):
        assert depo.rapor_klasoru(kotu) is None
    # eski sürüm temizliği bir kez çalışır
    prg = tmp_path / "prg"; (prg / "luca_kayitlari" / "k").mkdir(parents=True)
    (tmp_path / "veri" / "raporlar").mkdir(); (tmp_path / "veri" / "toplu" / "202609_x").mkdir(parents=True)
    (tmp_path / "veri" / "ayarlar.json").write_text("{}")
    assert depo.eski_surumu_temizle(prg) is True
    assert not (prg / "luca_kayitlari").exists() and not (tmp_path / "veri" / "raporlar").exists()
    assert not (tmp_path / "veri" / "toplu").exists() and (tmp_path / "veri" / "ayarlar.json").exists()
    assert depo.eski_surumu_temizle(prg) is False


def test_elle_giris_pencere_kapanirsa_devam(sahte_luca, tmp_path, monkeypatch):
    """Kullanıcı elle girer, paketi kendisi açar; giriş penceresi kendini kapatır → program çökmeden devam eder
    ve paketi ikinci kez açmaz."""
    monkeypatch.setenv("FK_HEADLESS", "1")
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fatura_kontrol.luca import otomasyon as O
    monkeypatch.setattr(O, "BASLANGIC", sahte_luca.replace("index.html", "giris2.html"))

    def kullanici(ctx, sayfa):
        with sayfa.expect_popup() as pp:
            sayfa.click("#sistemGiris")
        g = pp.value
        g.fill("#j_uye", "1"); g.fill("#j_kul", "a"); g.fill("#j_par", "b"); g.click("#girisBtn")
        g.fill("#guvenlikKodu", "7K3P"); g.press("#guvenlikKodu", "Enter")
        g.wait_for_url("**/main.erp")
        g.click("text=LUCA MALİ MÜŞAVİR PAKETİ")

    with O.oturum(tmp_path / "profil", kullanici) as (ctx, L):
        assert O.firma_listesi(L)
        assert sum("luca.do" in p.url for p in ctx.pages if not p.is_closed()) == 1


def test_luca_penceresi_acik_kalir(sahte_luca, tmp_path, monkeypatch):
    """Sorgu bitince Luca tarayıcısı kapanmaz; ikinci sorgu aynı oturuma bağlanır, yeniden giriş istemez."""
    import glob
    import socket
    import subprocess
    exe = (glob.glob("/opt/pw-browsers/chromium-*/chrome-linux/chrome") or [None])[0]
    if not exe:
        pytest.skip("chromium yok")
    s = socket.socket(); s.bind(("127.0.0.1", 0)); port = s.getsockname()[1]; s.close()
    monkeypatch.setenv("FK_CDP", "1")
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    monkeypatch.setenv("FK_TARAYICI_EXE", exe)
    monkeypatch.setenv("FK_TARAYICI_EK", "--headless=new --no-sandbox")
    from fatura_kontrol.luca import otomasyon as O
    monkeypatch.setattr(O, "CDP_PORT", port)
    monkeypatch.setattr(O, "BASLANGIC", sahte_luca)
    girisler = []

    def giris(ctx, sayfa):
        girisler.append(1)
        with sayfa.expect_popup():
            sayfa.click("#giris")
    try:
        with O.oturum(tmp_path / "profil", giris) as (ctx, L):
            ilk = O.firma_listesi(L)
        assert O._cdp_acik()                                      # tarayıcı açık kaldı
        with O.oturum(tmp_path / "profil", giris) as (ctx, L):   # ikinci sorgu
            ikinci = O.firma_listesi(L)
        assert girisler == [1] and ilk == ikinci                  # yeniden giriş istenmedi
    finally:
        subprocess.run(["pkill", "-f", f"remote-debugging-port={port}"])


def test_dovizli_fatura_kontrol_disi():
    """Şimdilik yalnız TL faturalar karşılaştırılır; dövizli fatura 'kontrol dışı' bilgisiyle listelenir
    ve muavindeki KDV satırı 'faturası bulunamadı' sayılmaz."""
    from decimal import Decimal as D
    from fatura_kontrol import eslestirme as E
    from fatura_kontrol.fatura_okuyucu import faturalari_oku
    from fatura_kontrol.muavin_okuyucu import MuavinSatiri
    import tempfile, zipfile
    x = ubl("EX02026000000077", "2026-09-12", FIRMA_VKN, "1234567890", "25", "5", "20").decode()
    x = x.replace("<cbc:DocumentCurrencyCode>TRY</cbc:DocumentCurrencyCode>",
                  "<cbc:DocumentCurrencyCode>USD</cbc:DocumentCurrencyCode>").replace('currencyID="TRY"', 'currencyID="USD"')
    x = x.replace("<cac:TaxTotal>", "<cac:PricingExchangeRate><cbc:CalculationRate>48.2238</cbc:CalculationRate>"
                  "</cac:PricingExchangeRate><cac:TaxTotal>", 1)
    with tempfile.TemporaryDirectory() as t:
        z = Path(t) / "a.zip"
        with zipfile.ZipFile(z, "w") as zf:
            zf.writestr("1.xml", x)
        ft, _ = faturalari_oku([z], FIRMA_VKN)
    assert ft[0].doviz == "USD" and ft[0].tl_kaynak.startswith("kur ile")   # kur var → KDV TL'ye çevrildi
    m = [MuavinSatiri(dosya="m.xlsx", satir=1, fis_tipi="Mahsup", hesap_kodu="391.10.20", hesap_adi="%20 Hesaplanan KDV", fis_tarihi=date(2026, 9, 30),
                      fis_no="5", aciklama="12/09/2026-EX02026000000077-ALICI", borc=D(0), alacak=D("241.12"),
                      belge_tarihi=date(2026, 9, 12), belge_no="EX02026000000077", unvan="ALICI")]
    r = E.esle(ft, m, donem=(date(2026, 9, 1), date(2026, 9, 30)))
    assert [s.durum for s in r.sonuclar] == [E.DOVIZ_TL_UYUMLU] and not r.faturasi_olmayanlar


def test_fatura_arsivi_tek_kopya(tmp_path, monkeypatch):
    """Luca ZIP'indeki başka aylara ait faturalar kendi aylarına bir kez yazılır; ham ZIP silinir."""
    import zipfile
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fatura_kontrol import arsiv, depo
    agu = depo.donem_klasoru("UKİ MADEN", date(2026, 8, 1)); agu.mkdir(parents=True)
    ham = agu / "gib_efatura_alis_1.zip"
    with zipfile.ZipFile(ham, "w") as z:
        z.writestr("a.xml", ubl("AAA2026000000001", "2026-08-10", "1234567890", FIRMA_VKN, "100", "20", "20"))
        z.writestr("b.xml", ubl("AAA2026000000002", "2026-09-05", "1234567890", FIRMA_VKN, "100", "20", "20"))
        z.writestr("b.html", b"<html>gorunum</html>")
    sonuc = arsiv.arsive_al("UKİ MADEN", [ham])
    eyl = depo.donem_klasoru("UKİ MADEN", date(2026, 9, 1))
    assert not ham.exists() and sonuc == {"2026-08": 1, "2026-09": 1}
    assert arsiv.mevcut_adlar(agu / "faturalar.zip") == {"gib_efatura_alis/AAA2026000000001.xml"}
    assert arsiv.mevcut_adlar(eyl / "faturalar.zip") == {"gib_efatura_alis/AAA2026000000002.xml"}
    # eylül sorgusu aynı faturayı yine getirirse ikinci kopya yazılmaz
    ham2 = eyl / "gib_efatura_alis_1.zip"
    with zipfile.ZipFile(ham2, "w") as z:
        z.writestr("b.xml", ubl("AAA2026000000002", "2026-09-05", "1234567890", FIRMA_VKN, "100", "20", "20"))
    assert arsiv.arsive_al("UKİ MADEN", [ham2]) == {"2026-09": 0}
    # Luca'dan alınan arşiv kontrolde yeniden yüklenince (gib_yukleme_dosya) ya da aynı fatura XML olarak eklenince
    # ikinci kopya yazılmaz, tür klasörü korunur (eski sürümde arşivde 2 kat fatura sayılıyordu: 572 / 286)
    import shutil
    yuk = eyl / "gib_yukleme_dosya_1.zip"; shutil.copyfile(eyl / "faturalar.zip", yuk)
    yuk2 = eyl / "gib_yukleme_dosya_2.zip"
    with zipfile.ZipFile(yuk2, "w") as z:
        z.writestr("x.xml", ubl("AAA2026000000002", "2026-09-05", "1234567890", FIRMA_VKN, "100", "20", "20"))
        z.writestr("y.xml", ubl("AAA2026000000003", "2026-09-06", "1234567890", FIRMA_VKN, "100", "20", "20"))
    assert arsiv.arsive_al("UKİ MADEN", [yuk, yuk2]) == {"2026-09": 1}
    assert arsiv.mevcut_adlar(eyl / "faturalar.zip") == {"gib_efatura_alis/AAA2026000000002.xml", "gib_yukleme_dosya/AAA2026000000003.xml"}
    arsiv.tur_sil(eyl, "gib_efatura_alis")
    assert arsiv.mevcut_adlar(eyl / "faturalar.zip") == {"gib_yukleme_dosya/AAA2026000000003.xml"}


def _usd(no, kdv_usd="5", ek="", kur=None, kdv_try=None):
    x = ubl(no, "2026-09-12", "1234567890", FIRMA_VKN, "25", kdv_usd, "20").decode()
    x = x.replace("<cbc:DocumentCurrencyCode>TRY</cbc:DocumentCurrencyCode>",
                  "<cbc:DocumentCurrencyCode>USD</cbc:DocumentCurrencyCode>").replace('currencyID="TRY"', 'currencyID="USD"')
    if kdv_try:      # KDV TL olarak da yazılmış (bazı entegratörler)
        x = x.replace(f'<cbc:TaxAmount currencyID="USD">{kdv_usd}</cbc:TaxAmount><cac:TaxSubtotal>',
                      f'<cbc:TaxAmount currencyID="TRY">{kdv_try}</cbc:TaxAmount><cac:TaxSubtotal>')
        x = x.replace(f'<cbc:TaxAmount currencyID="USD">{kdv_usd}</cbc:TaxAmount><cbc:Percent>',
                      f'<cbc:TaxAmount currencyID="TRY">{kdv_try}</cbc:TaxAmount><cbc:Percent>')
    if kur:
        x = x.replace("<cac:TaxTotal>", f"<cac:PricingExchangeRate><cbc:CalculationRate>{kur}</cbc:CalculationRate>"
                      "</cac:PricingExchangeRate><cac:TaxTotal>", 1)
    if ek:
        x = x.replace("<cbc:IssueDate>", f"<cbc:Note>{ek}</cbc:Note><cbc:IssueDate>", 1)
    from fatura_kontrol.fatura_okuyucu import xml_oku
    return xml_oku(x.encode(), "t.xml")


def test_dovizli_fatura_tl_karsiliklari():
    from decimal import Decimal as D
    from fatura_kontrol import eslestirme as E
    from fatura_kontrol.muavin_okuyucu import MuavinSatiri
    # 1) KDV TL olarak yazılmış (GNV/CWC tipi)
    f1 = _usd("GNV2026000004750", kdv_try="241.12")
    assert f1.tl_kaynak == "TL tutar faturada yazılı" and f1.kdv_toplam == D("241.12")
    # 2) TL KDV notta yazılı
    f2 = _usd("CWC2026000003686", ek="Kur: 48,2238 TL KARŞILIKLARI: Mal Hizmet 1.205,60 TL; KDV TL Karşılığı: 241,12 TL; "
                                     "KDV Dahil Toplam 1.446,72 TL")
    assert f2.tl_kaynak.startswith("TL tutar faturada yazılı") and f2.kdv_toplam == D("241.12")
    assert f2.kdv_oranlari == {D(20): D("241.12")}
    # 3) TL yok, yalnız kur (GEM tipi): döviz KDV × kur
    f3 = _usd("GEM2026000001680", kur="48.2238")
    assert f3.tl_kaynak.startswith("kur ile") and f3.kdv_toplam == D("241.12")
    # 4) hiçbiri yok → kontrol dışı
    f4 = _usd("XYZ2026000000001")
    assert f4.tl_kaynak is None
    for f in (f1, f2, f3, f4):
        f.yon = "ALIS"
    m = [MuavinSatiri(dosya="m", satir=1, fis_tipi="M", hesap_kodu="191.10.20", hesap_adi="%20 İndirilecek KDV",
                      fis_tarihi=date(2026, 9, 30), fis_no="5", aciklama="", borc=D("241.12"), alacak=D(0),
                      belge_tarihi=date(2026, 9, 12), belge_no=no, unvan="X")
         for no in ("GNV2026000004750", "CWC2026000003686", "GEM2026000001680")]
    r = E.esle([f1, f2, f3, f4], m, donem=(date(2026, 9, 1), date(2026, 9, 30)))
    d = {s.fatura.belge_no: s.durum for s in r.sonuclar}
    assert d == {"GNV2026000004750": E.DOVIZ_TL_UYUMLU, "CWC2026000003686": E.DOVIZ_TL_UYUMLU,
                 "GEM2026000001680": E.DOVIZ_TL_UYUMLU, "XYZ2026000000001": E.DOVIZLI}


def test_muavin_menuden_ve_oturum_dusunce_yeniden_giris(sahte_luca, tmp_path, monkeypatch):
    """Muavin Muhasebe > Raporlar > Defterler > Tüm Yazıcılar > Muavin Defter menüsüyle açılır (sağ paneldeki /
    Nokta Vuruşlu altındaki aynı adlı öğeye basılmaz). Oturum düşüp 'www.luca.com.tr adresinden giriş yapınız'
    sayfası gelirse o sekme kapatılır, açık giriş sayfasından Mali Müşavir Paketi'ne yeniden girilir ve muavin
    yeniden indirilir."""
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from playwright.sync_api import sync_playwright
    from fatura_kontrol.luca import otomasyon as O
    monkeypatch.setattr(O, "BASLANGIC", sahte_luca)
    monkeypatch.setattr(O, "GIRIS_BEKLEME", 60)
    kok = sahte_luca.rsplit("/", 1)[0]
    with sync_playwright() as pw:
        b = pw.chromium.launch(headless=True)
        ctx = b.new_context(accept_downloads=True)
        giris = ctx.new_page()
        giris.goto(kok + "/LUCASSO/main.erp?acik")
        with ctx.expect_page() as yeni:
            giris.get_by_text("LUCA MALİ MÜŞAVİR PAKETİ").click()
        L = yeni.value
        O.frame(L, "frm4", 30).wait_for_selector("#SirketCombo")
        L.evaluate("localStorage.setItem('dusur', '1')")           # ilk muavin açılışında oturum düşecek

        s = O.firma_isle(L, "id:110198728", date(2026, 9, 1), date(2026, 9, 30), [], tmp_path / "cikti")
        adimlar = {a["adim"]: a for a in s.adimlar}
        print(json.dumps(s.adimlar, ensure_ascii=False, indent=1))
        assert s.muavin and s.muavin.exists()
        assert adimlar["Muavin indirildi"]["durum"] == "tamam" and adimlar["Luca oturumu düştü"]
        assert s.L is not L and "luca.do" in s.L.url                # yeni Luca sekmesiyle devam edildi
        assert not any(O.luca_hatasi(p) for p in ctx.pages)         # hata sekmesi kapatıldı
        assert s.L.evaluate("sessionStorage.getItem('yanlis')") is None
        b.close()


@pytest.mark.parametrize("pager,beklenen", [
    ('<button onclick="sayfa(+1)">></button>', "sayfa(+1)"),
    ('<div class="sayfalama">Sayfa <select onchange="tik=\'sec\'+this.value"><option>1</option><option>2</option><option>3</option></select></div>'
     '<select onchange="tik=\'boy\'"><option>50</option><option selected>500</option></select>', "sayfa listesi"),
    ('<div class="pagination"><span class="aktif">1</span><a href="#" onclick="tik=\'n2\';return false">2</a><a href="#" onclick="tik=\'n3\'">3</a></div>', "sayfa 2"),
    ('<ul class="pager"><li><a onclick="tik=\'ok\'">»</a></li></ul>', "ileri"),
    ('<div class="pagination"><a class="disabled" onclick="tik=\'x\'">»</a></div><table><tr><td>2</td></tr></table>', ""),
])
def test_luca_sonraki_sayfa(pager, beklenen):
    """Luca e-belge listesi 500'er belge: sol alttaki sayfa değiştirme alanının türü ne olursa olsun sonraki sayfaya geçilir;
    sayfa kalmadıysa (ya da yalnız kapalı ok varsa) hiçbir şeye basılmaz. Listedeki '2' gibi hücrelere tıklanmaz."""
    from playwright.sync_api import sync_playwright
    from fatura_kontrol.luca import otomasyon as O
    with sync_playwright() as p:
        b = p.chromium.launch(); pg = b.new_page()
        pg.set_content(f"<script>var tik='';function sayfa(d){{tik='arti'+d}}</script><div>1 - 500 / 1.234 kayıt</div>{pager}")
        assert pg.evaluate(O.SONRAKI_SAYFA_JS, 1) == beklenen
        assert (pg.evaluate("tik") != "") == bool(beklenen) and pg.evaluate("tik") not in ("boy", "x")
        assert pg.evaluate(O.TOPLAM_JS) == 1234
        b.close()


def test_muhawin_oturumu_aktarilir(sahte_luca, tmp_path, monkeypatch):
    """MuhaWin'deki Luca oturumu (çerezler) otomasyon tarayıcısına aktarılır: giriş / doğrulama kodu istenmeden ürün
    seçim sayfasından Mali Müşavir Paketi kendiliğinden açılır; çerez dosyası okunur okunmaz silinir."""
    monkeypatch.setenv("FK_HEADLESS", "1")
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fatura_kontrol.luca import otomasyon as O
    kok = sahte_luca.rsplit("/", 1)[0]
    y = tmp_path / "c.json"
    y.write_text(json.dumps([{"name": "LUCASSO", "value": "abc", "url": kok + "/"}]), encoding="utf-8")
    monkeypatch.setenv("FK_LUCA_CEREZ", str(y))
    monkeypatch.setattr(O, "_CEREZ", None)
    monkeypatch.setattr(O, "SSO_ANA", kok + "/LUCASSO/main.erp?acik")
    monkeypatch.setattr(O, "BASLANGIC", sahte_luca)
    monkeypatch.setattr(O, "GIRIS_BEKLEME", 40)
    monkeypatch.setattr(O, "_OTO", {"paket": False})
    girisler = []
    with O.oturum(tmp_path / "profil", lambda ctx, s: girisler.append(1)) as (ctx, L):
        assert "luca.do" in L.url and O.firma_listesi(L)
        assert any(c["name"] == "LUCASSO" and c["value"] == "abc" for c in ctx.cookies())
    assert not y.exists() and "FK_LUCA_CEREZ" not in os.environ
    assert girisler == []                                       # giriş sayfası açılmadı


def test_muhawin_oturumu_gecersizse_normal_giris(sahte_luca, tmp_path, monkeypatch):
    monkeypatch.setenv("FK_HEADLESS", "1")
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fatura_kontrol.luca import otomasyon as O
    kok = sahte_luca.rsplit("/", 1)[0]
    monkeypatch.setattr(O, "_CEREZ", [{"name": "X", "value": "1", "url": kok + "/"}])
    monkeypatch.setattr(O, "SSO_ANA", kok + "/LUCASSO/giris.erp")       # oturum düşmüş: giriş sayfasına yönlenir
    monkeypatch.setattr(O, "BASLANGIC", sahte_luca)
    monkeypatch.setattr(O, "GIRIS_BEKLEME", 40)

    def giris(ctx, sayfa):
        with sayfa.expect_popup():
            sayfa.click("#giris")
    with O.oturum(tmp_path / "profil", giris) as (ctx, L):
        assert "luca.do" in L.url


def test_luca_oturumu_api(tmp_path, monkeypatch):
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fastapi.testclient import TestClient
    from fatura_kontrol.web import app as A, toplu_api
    c = TestClient(A.app)
    r = c.post("/api/toplu/luca-oturum", json={"cerezler": [
        {"name": "JSESSIONID", "value": "s1", "domain": "auygs.luca.com.tr", "path": "/Luca", "httpOnly": True, "secure": True, "fazla": 1},
        {"name": "SSO", "value": "s2", "url": "https://agiris.luca.com.tr/"},
        {"name": "kotu", "value": "x", "domain": ".ornek.com"}, {"name": "kotu2", "value": "x", "url": "https://luca.com.tr.ornek.com/"}]})
    assert r.json() == {"tamam": True, "adet": 2} and "s1" not in r.text
    assert "fazla" not in toplu_api._LUCA_OTURUM["cerez"][0]
    yakala = {}

    class P:
        def __init__(self, args, env=None, **k):
            yakala["env"] = env; y = env.get("FK_LUCA_CEREZ")
            yakala["icerik"] = json.loads(Path(y).read_text()) if y else None
            yakala["izin"] = oct(os.stat(y).st_mode & 0o777) if y else None

        def poll(self):
            return 0
    monkeypatch.setattr(toplu_api.subprocess, "Popen", P)
    toplu_api._is.clear()
    toplu_api._baslat("toplu", ["--ay", "2026-09"], tmp_path / "k")
    assert [x["name"] for x in yakala["icerik"]] == ["JSESSIONID", "SSO"]
    if os.name != "nt":
        assert yakala["izin"] == "0o600"
    os.remove(yakala["env"]["FK_LUCA_CEREZ"])
    c.post("/api/toplu/luca-oturum", json={"cerezler": []})            # Luca'dan çıkıldı
    toplu_api._baslat("toplu", ["--ay", "2026-09"], tmp_path / "k")
    assert "FK_LUCA_CEREZ" not in yakala["env"]
    toplu_api._is.clear()


def test_captcha_programda(sahte_luca, tmp_path, monkeypatch):
    """Doğrulama kodu programda: resim kaydedilir, kullanıcının programda yazdığı kod Luca'ya yazılır; yanlışsa yeni resim
    istenir; doğru kodla giriş tamamlanır ve Mali Müşavir Paketi'ne basılır."""
    import threading
    monkeypatch.setenv("FK_HEADLESS", "1")
    monkeypatch.setenv("FK_VERI", str(tmp_path / "veri"))
    from fatura_kontrol.luca import otomasyon as O
    from fatura_kontrol import ayarlar
    monkeypatch.setattr(O, "BASLANGIC", sahte_luca.replace("index.html", "giris2.html"))
    monkeypatch.setattr(O, "GIRIS_BEKLEME", 60)
    monkeypatch.setattr(O, "_CEREZ", [])
    monkeypatch.setattr(O, "_OTO", {"paket": False})
    monkeypatch.setattr(ayarlar, "luca_girisi", lambda: {"uye_no": "1212207", "kullanici": "ali", "parola": "x"})
    k = tmp_path / "cikti"; k.mkdir()
    monkeypatch.setattr(O, "CIKTI", k)
    olaylar, verilen = [], []
    monkeypatch.setattr(O, "DINLEYICI", lambda o, a: olaylar.append((o, a)))

    import time

    def kullanici():
        for kod in ("XXXX", "7K3P"):
            for _ in range(900):
                if (k / "captcha.png").exists():
                    break
                time.sleep(0.1)
            (k / "captcha.png").unlink(); verilen.append(kod)
            (k / "captcha_yanit.txt").write_text(kod, encoding="utf-8")
            time.sleep(0.5)
    t = threading.Thread(target=kullanici, daemon=True); t.start()
    with O.oturum(tmp_path / "profil") as (ctx, L):
        assert "luca.do" in L.url and O.firma_listesi(L)
    t.join(5)
    assert verilen == ["XXXX", "7K3P"]
    assert any(o == "captcha" and "kabul edilmedi" in a for o, a in olaylar)
    assert not (k / "captcha.png").exists() and not (k / "captcha_yanit.txt").exists()
