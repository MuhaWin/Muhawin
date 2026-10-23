"""test/ornek/ örnek dosyalarının üreticisi (test/e2e.js kullanır). Tüm içerik uydurmadır; bkz. BENIOKU.md.

    python3 test/ornek/uret.py          # Pillow, openpyxl, xlwt gerekir (fk/.venv'de: pip install pillow openpyxl xlwt)

Üretilen dosyalar depoya eklenir (testler ağsız ve hızlı kalsın). Rastgelelik sabit tohumla: her çalıştırmada aynı içerik.
"""
from __future__ import annotations

import random
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import xlwt
from PIL import Image, ImageDraw, ImageFont

KLASOR = Path(__file__).resolve().parent
YAZI = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"


# ---------------------------------------------------------------- ÖKC fiş fotoğrafları (e2e: "gerçek fiş fotoğrafı okundu")
def fis(ad: str, satirlar: list[str]) -> None:
    """Termal fiş benzeri görüntü: beyaz zemin, eşit aralıklı yazı; OCR'ın (tesseract, Türkçe) okuyacağı netlikte."""
    f = ImageFont.truetype(YAZI, 24)
    img = Image.new("L", (600, 50 + len(satirlar) * 34), 255)
    d = ImageDraw.Draw(img)
    for i, s in enumerate(satirlar):
        d.text((14, 20 + i * 34), s, font=f, fill=0)
    img.convert("RGB").save(KLASOR / ad, "JPEG", quality=92)


def fisler() -> None:
    # okc_fis1.jpg → 2026-10-15 | 0004 | 1000 | %20: 833,33 / 166,67 | kredi kartı
    fis("okc_fis1.jpg", [
        "     UYDURMA KIRTASİYE LTD. ŞTİ.",
        "     ÖRNEK CAD. NO:5 ANTALYA",
        "   DENEME V.D. 1111111114",
        "",
        "TARİH : 15/10/2026     SAAT : 10:05",
        "FİŞ NO : 0004",
        "",
        "TONER                  %20   *600,00",
        "KAĞIT                  %20   *400,00",
        "------------------------------------",
        "TOPKDV                       *166,67",
        "TOPLAM                     *1.000,00",
        "KREDİ KARTI                *1.000,00",
        "",
        "Z NO : 0101    EKÜ NO : 0001",
    ])
    # okc_fis2.jpg → 2026-10-17 | 0007 | 1060 | %10: 963,64 / 96,36 | kredi kartı; kalem "YEMEK" → 770 yemek gideri
    fis("okc_fis2.jpg", [
        "      DEMO LOKANTA LTD. ŞTİ.",
        "     SAHTE SOK. NO:7 ANTALYA",
        "   DENEME V.D. 2222222220",
        "",
        "TARİH : 17/10/2026     SAAT : 13:20",
        "FİŞ NO : 0007",
        "",
        "YEMEK                  %10   *800,00",
        "İÇECEK                 %10   *260,00",
        "------------------------------------",
        "TOPKDV                        *96,36",
        "TOPLAM                     *1.060,00",
        "KREDİ KARTI                *1.060,00",
        "",
        "Z NO : 0102    EKÜ NO : 0001",
    ])


# ---------------------------------------------------------------- mizanlar (e2e: Çek / Senet; HESAP KODU başlığı, r[0] kod, r[1] ad, r[3] borç, r[4] alacak)
INKA_MUSTERI = [("120.01.001", "SANTUR TURİZM A.Ş."), ("120.01.002", "ÖRNEK GIDA LTD. ŞTİ."), ("120.01.003", "DEMO TEKSTİL A.Ş.")]
INKA_SATICI = [("320.01.001", "UYDURMA İNŞAAT LTD. ŞTİ."), ("320.01.002", "TEST MOBİLYA A.Ş.")]
SAVOR_MUSTERI = ["MAVİ DENİZ OTELCİLİK A.Ş.", "YEŞİL VADİ GIDA LTD. ŞTİ.", "KUZEY YILDIZI TEKSTİL A.Ş.", "GÜNEY RÜZGARI TURİZM LTD. ŞTİ.",
                 "BEYAZ KUM İNŞAAT A.Ş.", "KIRMIZI ELMA PAZARLAMA LTD. ŞTİ.", "MOR DAĞ OTOMOTİV A.Ş.", "TURUNCU GÜNEŞ ENERJİ LTD. ŞTİ."]
SAVOR_SATICI = ["ÖRNEK AMBALAJ SANAYİ A.Ş.", "DENEME LOJİSTİK LTD. ŞTİ.", "SAHTE KİMYA A.Ş.", "UYDURMA MAKİNE LTD. ŞTİ."]
AY = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran", "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]


def mizan(ad: str, firma: str, hesaplar: list[tuple[str, str]]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([f"{firma} – MİZAN (01.01.2026 - 30.09.2026) – örnek, uydurma veri"])
    ws.append([])
    ws.append(["HESAP KODU", "HESAP ADI", "", "BORÇ", "ALACAK"])
    for kod, hadi in hesaplar:
        ws.append([kod, hadi, "", 0, 0])
    wb.save(KLASOR / ad)


def vade_hesaplari(ana: str, ad: str, aylar: list[tuple[int, int]]) -> list[tuple[str, str]]:
    out = [(ana, ad)]
    for y in sorted({y for y, _ in aylar}):
        out.append((f"{ana}.{y % 100:02d}", f"{y} Yılı {ad}"))
    out += [(f"{ana}.{y % 100:02d}.{m:02d}", f"{AY[m - 1]} {y} {ad}") for y, m in aylar]
    return out


def mizanlar() -> None:
    ortak = [("100", "KASA"), ("100.01", "MERKEZ KASA"), ("102", "BANKALAR"), ("102.01", "GARANTİ BANKASI TL"), ("102.02", "İŞ BANKASI TL"),
             ("102.03", "ZİRAAT BANKASI TL")]
    # İNKA (Mikro bordroları): 101 vade hesapları var; 103 (verilen çek) vade hesapları YOK → e2e açılmasını ister
    mizan("mizan_inka.xlsx", "İNKA ÖRNEK A.Ş.", ortak
          + vade_hesaplari("101", "Vadeli Alınan Çekler", [(2026, 10), (2026, 11), (2026, 12)])
          + [("103", "VERİLEN ÇEKLER VE ÖDEME EMİRLERİ")]
          + [("120", "ALICILAR"), ("120.01", "YURTİÇİ ALICILAR")] + INKA_MUSTERI
          + [("320", "SATICILAR"), ("320.01", "YURTİÇİ SATICILAR")] + INKA_SATICI)
    # SAVOR (Logo listeleri): tüm vade hesapları var (Logo eklenince hesap açma sorusu çıkmaz)
    aylar = [(2026, 10), (2026, 11), (2026, 12), (2027, 1), (2027, 2), (2027, 3)]
    mizan("mizan_savor.xlsx", "SAVOR ÖRNEK LTD. ŞTİ.", ortak
          + vade_hesaplari("101", "Vadeli Alınan Çekler", aylar)
          + vade_hesaplari("103", "Vadeli Verilen Çekler", aylar)
          + [("120", "ALICILAR"), ("120.01", "YURTİÇİ ALICILAR")] + [(f"120.01.{i:03d}", a) for i, a in enumerate(SAVOR_MUSTERI, 1)]
          + [("320", "SATICILAR"), ("320.01", "YURTİÇİ SATICILAR")] + [(f"320.01.{i:03d}", a) for i, a in enumerate(SAVOR_SATICI, 1)])


# ---------------------------------------------------------------- Mikro çek bordroları (.xls; e2e: Mikro giriş / çıkış)
MIKRO_BASLIK = ["Evrak No", "Evrak tarihi", "Borc Alacak", "Cari kodu", "Cari adı", "Cinsi", "Referans", "Vade",
                "Karşı cari kodu", "Karşı cari adı", "Detay - Açıklama", "Meblağ"]


def mikro_xls(ad: str, bordrolar: list[dict]) -> None:
    """Mikro › Çek/senet bordro dökümü: bordro başlık satırı (Evrak No dolu) ve altında evrak satırları (Cinsi, Vade, Meblağ)."""
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Bordro")
    ws.write(0, 0, "ÇEK / SENET BORDRO DÖKÜMÜ – örnek, uydurma veri")
    for j, b in enumerate(MIKRO_BASLIK):
        ws.write(2, j, b)
    i = 3
    for b in bordrolar:
        for j, v in enumerate([b["evrak"], b["tarih"], b["ba"], b["kod"], b["ad"]]):
            ws.write(i, j, v)
        i += 1
        for c in b["cekler"]:
            satir = ["", "", "", "", "", c["cins"], c["ref"], c["vade"], c.get("kk", ""), c.get("ka", ""), f"{c['no']} / {c['banka']}", c["tutar"]]
            for j, v in enumerate(satir):
                if v != "":
                    ws.write(i, j, v)
            i += 1
    wb.save(str(KLASOR / ad))


def mikro() -> None:
    mc = "Müşteri çeki"
    mikro_xls("mikro_giris.xls", [
        {"evrak": "CG-0001", "tarih": "05.09.2026", "ba": "Alacak", "kod": "120.01.001", "ad": "SANTUR TURİZM A.Ş.", "cekler": [
            {"cins": mc, "ref": "R1001", "vade": "15.10.2026", "no": "1000101", "banka": "GARANTİ BANKASI", "tutar": 15000.00},
            {"cins": mc, "ref": "R1002", "vade": "20.11.2026", "no": "1000102", "banka": "GARANTİ BANKASI", "tutar": 22500.00}]},
        {"evrak": "CG-0002", "tarih": "12.09.2026", "ba": "Alacak", "kod": "120.01.002", "ad": "ÖRNEK GIDA LTD. ŞTİ.", "cekler": [
            {"cins": mc, "ref": "R2001", "vade": "30.10.2026", "no": "2000201", "banka": "İŞ BANKASI", "tutar": 8750.50},
            {"cins": mc, "ref": "R2002", "vade": "10.12.2026", "no": "2000202", "banka": "İŞ BANKASI", "tutar": 12000.00}]},
        {"evrak": "CG-0003", "tarih": "20.09.2026", "ba": "Alacak", "kod": "120.01.003", "ad": "DEMO TEKSTİL A.Ş.", "cekler": [
            {"cins": mc, "ref": "R3001", "vade": "05.11.2026", "no": "3000301", "banka": "ZİRAAT BANKASI", "tutar": 6400.00}]},
    ])
    mikro_xls("mikro_cikis.xls", [
        # müşteri çeki satıcıya ciro (cari 320, Borç) → olay "ciro"
        {"evrak": "CC-0001", "tarih": "25.09.2026", "ba": "Borç", "kod": "320.01.001", "ad": "UYDURMA İNŞAAT LTD. ŞTİ.", "cekler": [
            {"cins": mc, "ref": "R1001", "vade": "15.10.2026", "no": "1000101", "banka": "GARANTİ BANKASI", "tutar": 15000.00},
            {"cins": mc, "ref": "R3001", "vade": "05.11.2026", "no": "3000301", "banka": "ZİRAAT BANKASI", "tutar": 6400.00}]},
        # kendi çekimiz satıcıya (Borç) → olay "verilen"; 103.26.11 / 103.26.12 mizanda yok
        {"evrak": "CC-0002", "tarih": "28.09.2026", "ba": "Borç", "kod": "320.01.002", "ad": "TEST MOBİLYA A.Ş.", "cekler": [
            {"cins": "Kendi çekimiz", "ref": "K0001", "vade": "30.11.2026", "kk": "102.01", "ka": "GARANTİ BANKASI", "no": "9000001", "banka": "GARANTİ BANKASI", "tutar": 18000.00},
            {"cins": "Kendi çekimiz", "ref": "K0002", "vade": "31.12.2026", "kk": "102.01", "ka": "GARANTİ BANKASI", "no": "9000002", "banka": "GARANTİ BANKASI", "tutar": 9500.00}]},
    ])


# ---------------------------------------------------------------- Logo çek / senet listeleri (.xlsx; e2e: Logo müşteri ve kendi)
LOGO_BASLIK = ["Portföy No.", "Seri No.", "Türü", "Vade", "Tutarı", "Durumu", "Borçlu", "Ciro Eden", "Ciro Edilen", "Çekin Bankası"]
BANKALAR = ["GARANTİ BANKASI", "İŞ BANKASI", "ZİRAAT BANKASI"]


def logo_xlsx(ad: str, baslik: str, satirlar: list[list]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([baslik])
    ws.append(["örnek, uydurma veri"])
    ws.append(LOGO_BASLIK)
    for s in satirlar:
        ws.append(s)
    wb.save(KLASOR / ad)


def logo() -> None:
    rnd = random.Random(2026)
    vade = lambda: (date(2026, 10, 1) + timedelta(days=rnd.randrange(181))).strftime("%d.%m.%Y")   # 10.2026–03.2027
    tutar = lambda: round(rnd.uniform(1000, 60000), 2)
    # müşteri çekleri: 260 çek; giriş + durumuna göre ciro / tahsile / tahsil / iade → listede 300'den çok hareket
    musteri = []
    for i in range(1, 261):
        k = i % 20
        durum = "Portföyde" if k < 8 else "Ciro Edildi" if k < 14 else "Tahsile Verildi" if k < 18 else "Tahsil Edildi" if k == 18 else "İade Edildi"
        v = vade()
        if durum == "Tahsil Edildi":                      # tahsil fişi vade tarihiyle kesilir: iki tarihte topla (fiş sayısı küçük kalsın)
            v = rnd.choice(["15.10.2026", "16.11.2026"])
        borclu = rnd.choice(SAVOR_MUSTERI)
        musteri.append([f"MP-{i:04d}", f"{4000000 + i}", "Müşteri Çeki", v, tutar(), durum, borclu, borclu,
                        rnd.choice(SAVOR_SATICI) if durum == "Ciro Edildi" else "", rnd.choice(BANKALAR)])
    logo_xlsx("logo_musteri.xlsx", "SAVOR ÖRNEK LTD. ŞTİ. – MÜŞTERİ ÇEK/SENET LİSTESİ", musteri)
    # kendi çeklerimiz: 60 çek; verildi, bir kısmı ödendi (iki vade tarihinde)
    kendi = []
    for i in range(1, 61):
        odendi = i % 10 == 0
        v = rnd.choice(["30.10.2026", "30.11.2026"]) if odendi else vade()
        kendi.append([f"KP-{i:04d}", f"{9000000 + i}", "Kendi Çeki", v, tutar(), "Ödendi" if odendi else "Verildi", "",
                      "", rnd.choice(SAVOR_SATICI), "GARANTİ BANKASI"])
    logo_xlsx("logo_kendi.xlsx", "SAVOR ÖRNEK LTD. ŞTİ. – KENDİ ÇEK/SENET LİSTESİ", kendi)


if __name__ == "__main__":
    fisler()
    mizanlar()
    mikro()
    logo()
    print("Yazıldı:", ", ".join(sorted(p.name for p in KLASOR.iterdir() if p.suffix in (".jpg", ".xlsx", ".xls"))))
